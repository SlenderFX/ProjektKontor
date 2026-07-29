from __future__ import annotations

import base64
import io
import json
import os
import tarfile
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from projektkontor.backup import create_backup
from projektkontor.config import Config
from projektkontor.server import App, HttpError

TEST_TEACHER_PASSWORD = "persoenliches-testkennwort"


class Client:
    def __init__(self, app: App, auto_privacy: bool = True):
        self.app = app
        self.auto_privacy = auto_privacy
        self.cookie = ""
        self.csrf = ""

    def request(self, method: str, path: str, body=None):
        raw = json.dumps(body or {}).encode() if body is not None else b""
        captured = {}

        def start_response(status, headers):
            captured["status"] = int(status.split()[0])
            captured["headers"] = dict(headers)

        environ = {
            "REQUEST_METHOD": method,
            "PATH_INFO": path,
            "QUERY_STRING": "",
            "CONTENT_LENGTH": str(len(raw)),
            "CONTENT_TYPE": "application/json",
            "HTTP_COOKIE": self.cookie,
            "HTTP_X_CSRF_TOKEN": self.csrf,
            "wsgi.input": io.BytesIO(raw),
        }
        payload = b"".join(self.app(environ, start_response))
        if "Set-Cookie" in captured["headers"]:
            self.cookie = captured["headers"]["Set-Cookie"].split(";", 1)[0]
        data = json.loads(payload) if captured["headers"].get("Content-Type", "").startswith("application/json") else payload
        if isinstance(data, dict) and data.get("csrf"):
            self.csrf = data["csrf"]
        if self.auto_privacy and isinstance(data, dict) and data.get("password_change_required"):
            return self.request("POST", "/api/initial-password", {
                "password_change_token": data["password_change_token"],
                "new_password": TEST_TEACHER_PASSWORD,
            })
        if self.auto_privacy and isinstance(data, dict) and data.get("privacy_required"):
            return self.request("POST", "/api/privacy/accept", {
                "privacy_token": data["privacy_token"], "privacy_version": data["privacy_version"]
            })
        return captured["status"], data, captured["headers"]


class AppFlowTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        root.joinpath("uploads").mkdir()
        root.joinpath("reports").mkdir()
        self.app = App(Config("127.0.0.1", 8080, root, 25 * 1024 * 1024))
        self.client = Client(self.app)

    def tearDown(self):
        self.temp.cleanup()

    def setup_teacher(self, client=None, username="lehrkraft_test"):
        client = client or self.client
        status, _, _ = client.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        self.assertEqual(status, 200)
        status, teacher, _ = client.request("POST", "/api/teachers", {
            "first_name": "Jeroen", "username": username
        })
        self.assertEqual(status, 200)
        status, _, _ = client.request("POST", "/api/licenses", {
            "customer_name": "Testlizenz", "organization": "Testschule",
            "plan": "single", "seat_limit": 1, "amount_cents": 7900,
            "payment_status": "paid", "status": "active",
            "starts_on": "2020-01-01", "ends_on": "2099-12-31",
            "teacher_ids": [teacher["id"]],
        })
        self.assertEqual(status, 200)
        client.request("POST", "/api/logout", {})
        status, _, _ = client.request("POST", "/api/login", {
            "login_type": "teacher", "username": username, "password": teacher["initial_password"]
        })
        self.assertEqual(status, 200)
        return teacher

    def test_teacher_class_students_project_and_task(self):
        self.setup_teacher()

        status, klass, _ = self.client.request("POST", "/api/classes", {"name": "GH24"})
        self.assertEqual(status, 200)
        class_id = klass["id"]

        status, imported, _ = self.client.request("POST", f"/api/classes/{class_id}/import", {
            "rows": [{"first_name": "Lena"}, {"first_name": "Lena"}, {"first_name": "Yusuf"}]
        })
        self.assertEqual(status, 200)
        self.assertEqual([row["username"] for row in imported["created"]], ["lena01", "lena02", "yusuf01"])
        lead_id = imported["created"][0]["id"]
        member_ids = [row["id"] for row in imported["created"]]

        status, project, _ = self.client.request("POST", "/api/projects", {
            "title": "Messeprojekt", "class_id": class_id, "description": "Eine Messe planen",
            "start_at": "2026-07-16T08:00", "end_at": "2026-07-23T12:00",
            "project_lead_id": lead_id, "member_ids": member_ids,
        })
        self.assertEqual(status, 200)
        project_id = project["id"]

        _, later_import, _ = self.client.request("POST", f"/api/classes/{class_id}/import", {
            "rows": [{"first_name": "Nora"}]
        })
        nora_id = later_import["created"][0]["id"]
        _, before_team, _ = self.client.request("GET", f"/api/projects/{project_id}")
        self.assertIn(nora_id, [student["id"] for student in before_team["class_students"]])
        self.assertNotIn(nora_id, [member["id"] for member in before_team["members"]])
        nora_before = next(student for student in before_team["class_students"] if student["id"] == nora_id)
        self.assertIsNone(nora_before["team_id"])

        status, team, _ = self.client.request("POST", f"/api/projects/{project_id}/teams", {
            "name": "Marketing", "responsibility": "Kommunikation", "color": "#2563EB",
            "member_ids": [lead_id], "lead_ids": [lead_id],
        })
        self.assertEqual(status, 200)
        team_id = team["id"]
        stored_member = self.app.db.one("SELECT is_lead FROM team_members WHERE team_id=? AND user_id=?", (team_id, lead_id))
        self.assertEqual(stored_member["is_lead"], 1)

        status, second_team, _ = self.client.request("POST", f"/api/projects/{project_id}/teams", {
            "name": "Vertrieb", "responsibility": "Kundschaft", "color": "#16A34A",
            "member_ids": [nora_id], "lead_ids": [nora_id],
        })
        self.assertEqual(status, 200)
        self.assertIsNotNone(self.app.db.one("SELECT 1 ok FROM project_members WHERE project_id=? AND user_id=?", (project_id, nora_id)))
        _, after_team, _ = self.client.request("GET", f"/api/projects/{project_id}")
        nora_after = next(student for student in after_team["class_students"] if student["id"] == nora_id)
        self.assertEqual(nora_after["team_id"], second_team["id"])
        self.assertEqual(nora_after["is_lead"], 1)

        status, task, _ = self.client.request("POST", f"/api/projects/{project_id}/tasks", {
            "title": "Flyer abstimmen", "team_ids": [team_id], "weight": 2,
            "start_at": "2026-07-16T09:00", "due_at": "2026-07-16T11:30"
        })
        self.assertEqual(status, 200)
        task_id = task["ids"][0]
        self.client.request("POST", f"/api/tasks/{task_id}/assignees", {"user_ids": [lead_id]})

        status, all_users, _ = self.client.request("GET", "/api/users")
        self.assertEqual(status, 200)
        lead_context = next(account for account in all_users if account["id"] == lead_id)
        self.assertEqual(lead_context["class_name"], "GH24")
        self.assertEqual(lead_context["projects"][0]["title"], "Messeprojekt")
        self.assertEqual(lead_context["teams"][0]["name"], "Marketing")
        self.assertEqual(lead_context["teams"][0]["is_lead"], 1)
        self.assertEqual(lead_context["open_task_count"], 1)
        self.assertTrue(lead_context["access_code"])

        status, detail, _ = self.client.request("GET", f"/api/projects/{project_id}")
        self.assertEqual(status, 200)
        self.assertEqual(detail["tasks"][0]["assignees"][0]["first_name"], "Lena")
        self.assertEqual(detail["tasks"][0]["due_at"], "2026-07-16T11:30")
        self.assertEqual(detail["progress"]["overall"], 0)

        status, duplicate_color, _ = self.client.request("POST", f"/api/projects/{project_id}/teams", {
            "name": "Einkauf", "responsibility": "Beschaffung", "color": "#2563EB"
        })
        self.assertEqual(status, 409)
        self.assertIn("bereits verwendet", duplicate_color["error"])

        status, duplicate_name, _ = self.client.request("POST", f"/api/projects/{project_id}/teams", {
            "name": "marketing", "responsibility": "Doppelt", "color": "#DC2626"
        })
        self.assertEqual(status, 409)
        self.assertIn("Teamname", duplicate_name["error"])

        status, comment, _ = self.client.request("POST", f"/api/tasks/{task_id}/comments", {"body": "Der Entwurf ist bereit."})
        self.assertEqual(status, 200)
        status, report, _ = self.client.request("POST", f"/api/projects/{project_id}/report", {})
        self.assertEqual(status, 200)
        report_row = self.app.db.one("SELECT * FROM reports WHERE id=?", (report["id"],))
        self.assertTrue((self.app.config.report_dir / report_row["stored_name"]).is_file())

        status, template, _ = self.client.request("POST", f"/api/projects/{project_id}/template", {"name": "Messevorlage"})
        self.assertEqual(status, 200)
        status, copied, _ = self.client.request("POST", "/api/projects", {
            "title": "Messeprojekt 2", "class_id": class_id, "description": "Kopie",
            "start_at": "2026-08-01T08:00", "end_at": "2026-08-08T12:00",
            "project_lead_id": lead_id, "member_ids": member_ids, "template_id": template["id"],
        })
        self.assertEqual(status, 200)
        copied_marketing = self.app.db.one("SELECT id,responsibility,color FROM teams WHERE project_id=? AND name='Marketing'", (copied["id"],))
        self.assertIsNotNone(copied_marketing)
        self.assertEqual(copied_marketing["responsibility"], "Kommunikation")
        self.assertEqual(copied_marketing["color"], "#2563EB")
        self.assertEqual(self.app.db.one("SELECT COUNT(*) count FROM team_members WHERE team_id=?", (copied_marketing["id"],))["count"], 0)
        self.assertIsNotNone(self.app.db.one("SELECT id FROM teams WHERE project_id=? AND name='Vertrieb'", (copied["id"],)))
        self.assertIsNotNone(self.app.db.one("SELECT id FROM tasks WHERE project_id=? AND title='Flyer abstimmen'", (copied["id"],)))

        status, archived, _ = self.client.request("POST", f"/api/projects/{project_id}/archive", {"confirm_cleanup": True})
        self.assertEqual(status, 200)
        self.assertIsNone(self.app.db.one("SELECT id FROM tasks WHERE project_id=?", (project_id,)))
        self.assertIsNotNone(self.app.db.one("SELECT id FROM reports WHERE project_id=?", (project_id,)))

    def test_student_codes_are_readable_to_teacher(self):
        self.setup_teacher()
        _, klass, _ = self.client.request("POST", "/api/classes", {"name": "KB25"})
        _, imported, _ = self.client.request("POST", f"/api/classes/{klass['id']}/import", {"rows": [{"first_name": "Mia"}]})
        _, users, _ = self.client.request("GET", f"/api/classes/{klass['id']}/users")
        self.assertEqual(users[0]["access_code"], imported["created"][0]["access_code"])

    def test_teacher_can_create_student_account_manually(self):
        self.setup_teacher()
        _, klass, _ = self.client.request("POST", "/api/classes", {"name": "KB26"})
        status, first, _ = self.client.request("POST", f"/api/classes/{klass['id']}/users", {"first_name": "Lena"})
        self.assertEqual(status, 200)
        self.assertEqual(first["username"], "lena01")
        self.assertEqual(len(first["access_code"]), 8)
        status, second, _ = self.client.request("POST", f"/api/classes/{klass['id']}/users", {
            "first_name": "Lena", "username": "einkauf01", "access_code": "KONTOR-2026"
        })
        self.assertEqual(status, 200)
        self.assertEqual(second["username"], "einkauf01")
        self.assertEqual(second["access_code"], "KONTOR-2026")
        _, users, _ = self.client.request("GET", f"/api/classes/{klass['id']}/users")
        self.assertEqual(len(users), 2)

        status, duplicate_class, _ = self.client.request("POST", "/api/classes", {"name": "kb26"})
        self.assertEqual(status, 409)
        self.assertEqual(duplicate_class["error"], "Diese Klasse existiert bereits.")

        _, target_class, _ = self.client.request("POST", "/api/classes", {"name": "KB27"})
        status, renamed_class, _ = self.client.request("PATCH", f"/api/classes/{target_class['id']}", {"name": "KB28"})
        self.assertEqual(status, 200)
        self.assertEqual(renamed_class["name"], "KB28")
        status, _, _ = self.client.request("PATCH", f"/api/users/{first['id']}", {"class_id": target_class["id"]})
        self.assertEqual(status, 200)
        _, original_users, _ = self.client.request("GET", f"/api/classes/{klass['id']}/users")
        _, target_users, _ = self.client.request("GET", f"/api/classes/{target_class['id']}/users")
        self.assertNotIn(first["id"], [account["id"] for account in original_users])
        self.assertIn(first["id"], [account["id"] for account in target_users])

    def test_student_can_login_with_generated_access_code(self):
        self.setup_teacher()
        _, klass, _ = self.client.request("POST", "/api/classes", {"name": "KB29"})
        _, account, _ = self.client.request("POST", f"/api/classes/{klass['id']}/users", {"first_name": "Markus"})
        status, _, _ = self.client.request("POST", "/api/logout", {})
        self.assertEqual(status, 200)
        status, login, login_headers = self.client.request("POST", "/api/login", {
            "username": account["username"].upper(),
            "password": f"  {account['access_code'].lower()}  ",
        })
        self.assertEqual(status, 200)
        self.assertTrue(login["ok"])
        self.assertIn("Max-Age=21600", login_headers["Set-Cookie"])
        status, bootstrap, _ = self.client.request("GET", "/api/bootstrap")
        self.assertEqual(status, 200)
        self.assertEqual(bootstrap["user"]["id"], account["id"])
        self.assertEqual(bootstrap["user"]["role"], "student")
        session = self.app.db.one("SELECT created_at,expires_at FROM sessions WHERE user_id=? ORDER BY id DESC LIMIT 1", (account["id"],))
        session_duration = __import__("datetime").datetime.fromisoformat(session["expires_at"]) - __import__("datetime").datetime.fromisoformat(session["created_at"])
        self.assertEqual(session_duration.total_seconds(), 6 * 60 * 60)

        second_device = Client(self.app)
        status, _, _ = second_device.request("POST", "/api/login", {
            "username": account["username"], "password": account["access_code"]
        })
        self.assertEqual(status, 200)
        self.assertEqual(self.app.db.one("SELECT COUNT(*) count FROM sessions WHERE user_id=?", (account["id"],))["count"], 1)
        status, first_device_bootstrap, _ = self.client.request("GET", "/api/bootstrap")
        self.assertEqual(status, 200)
        self.assertIsNone(first_device_bootstrap["user"])
        status, second_device_bootstrap, _ = second_device.request("GET", "/api/bootstrap")
        self.assertEqual(status, 200)
        self.assertEqual(second_device_bootstrap["user"]["id"], account["id"])

    def test_student_roles_deadline_requests_and_self_organization(self):
        teacher = self.client
        self.setup_teacher(teacher)
        _, klass, _ = teacher.request("POST", "/api/classes", {"name": "WB30"})
        _, imported, _ = teacher.request("POST", f"/api/classes/{klass['id']}/import", {
            "rows": [{"first_name": "Markus"}, {"first_name": "Lena"}]
        })
        markus, lena = imported["created"]
        _, project, _ = teacher.request("POST", "/api/projects", {
            "title": "Rollenprojekt", "class_id": klass["id"], "has_start": False, "has_end": False,
            "project_lead_id": lena["id"], "member_ids": [markus["id"], lena["id"]],
        })
        _, team, _ = teacher.request("POST", f"/api/projects/{project['id']}/teams", {
            "name": "Einkauf", "color": "#2563EB", "member_ids": [markus["id"]], "lead_ids": []
        })
        _, task, _ = teacher.request("POST", f"/api/projects/{project['id']}/tasks", {
            "title": "Angebote prüfen", "team_ids": [team["id"]], "due_at": "2026-08-01"
        })
        task_id = task["ids"][0]
        teacher.request("POST", f"/api/tasks/{task_id}/assignees", {"user_ids": [markus["id"]]})

        student = Client(self.app)
        status, _, _ = student.request("POST", "/api/login", {
            "username": markus["username"], "password": markus["access_code"]
        })
        self.assertEqual(status, 200)
        status, forbidden_team, _ = student.request("POST", f"/api/projects/{project['id']}/teams", {
            "name": "Logistik", "color": "#16A34A", "member_ids": [], "lead_ids": []
        })
        self.assertEqual(status, 403)
        self.assertIn("deaktiviert", forbidden_team["error"])
        status, _, _ = student.request("POST", f"/api/tasks/{task_id}/assignees", {"user_ids": [markus["id"]]})
        self.assertEqual(status, 403)

        teacher.request("PATCH", f"/api/projects/{project['id']}", {"allow_student_organization": True})
        status, proposed_team, _ = student.request("POST", f"/api/projects/{project['id']}/teams", {
            "name": "Logistik", "color": "#16A34A", "member_ids": [], "lead_ids": []
        })
        self.assertEqual(status, 200)
        self.assertEqual(proposed_team["status"], "proposed")

        status, request, _ = student.request("POST", f"/api/tasks/{task_id}/deadline-requests", {
            "requested_due_at": "2026-08-05T12:00", "reason": "Lieferant antwortet später"
        })
        self.assertEqual(status, 200)
        status, duplicate_request, _ = student.request("POST", f"/api/tasks/{task_id}/deadline-requests", {
            "requested_due_at": "2026-08-06T12:00", "reason": "Noch später"
        })
        self.assertEqual(status, 409)
        status, _, _ = teacher.request("PATCH", f"/api/deadline-requests/{request['id']}", {"decision": "approved"})
        self.assertEqual(status, 200)
        self.assertEqual(self.app.db.one("SELECT due_at FROM tasks WHERE id=?", (task_id,))["due_at"], "2026-08-05T12:00")

        _, comment, _ = student.request("POST", f"/api/tasks/{task_id}/comments", {"body": "Erste Fassung"})
        status, _, _ = student.request("PATCH", f"/api/comments/{comment['id']}", {"body": "Überarbeitete Fassung"})
        self.assertEqual(status, 200)
        stored = self.app.db.one("SELECT body,edited_at FROM comments WHERE id=?", (comment["id"],))
        self.assertEqual(stored["body"], "Überarbeitete Fassung")
        self.assertTrue(stored["edited_at"])
        status, _, _ = teacher.request("DELETE", f"/api/comments/{comment['id']}", {})
        self.assertEqual(status, 200)

        pdf_data = base64.b64encode(b"%PDF-test-project-file").decode()
        status, project_upload, _ = teacher.request("POST", f"/api/projects/{project['id']}/uploads", {
            "name": "Projektauftrag.pdf", "media_type": "application/pdf", "file_base64": pdf_data
        })
        self.assertEqual(status, 200)
        status, team_upload, _ = student.request("POST", f"/api/projects/{project['id']}/uploads", {
            "name": "Teamdatei.pdf", "media_type": "application/pdf", "file_base64": pdf_data,
            "team_id": team["id"],
        })
        self.assertEqual(status, 200)
        self.assertNotEqual(project_upload["id"], team_upload["id"])

        _, outsider, _ = teacher.request("POST", f"/api/classes/{klass['id']}/users", {"first_name": "Nora"})
        outsider_client = Client(self.app)
        outsider_client.request("POST", "/api/login", {"username": outsider["username"], "password": outsider["access_code"]})
        status, _, _ = outsider_client.request("GET", f"/api/projects/{project['id']}")
        self.assertEqual(status, 403)

    def test_excel_duplicate_rows_require_confirmation(self):
        self.setup_teacher()
        _, klass, _ = self.client.request("POST", "/api/classes", {"name": "WB31"})
        rows = [{"first_name": "Markus", "duplicate": True}]
        status, result, _ = self.client.request("POST", f"/api/classes/{klass['id']}/import", {"rows": rows})
        self.assertEqual(status, 409)
        self.assertIn("Namensdubletten", result["error"])
        status, result, _ = self.client.request("POST", f"/api/classes/{klass['id']}/import", {
            "rows": rows, "confirm_duplicates": True
        })
        self.assertEqual(status, 200)
        self.assertEqual(len(result["created"]), 1)

    def test_excel_import_can_create_and_assign_multiple_classes(self):
        self.setup_teacher()
        _, klass, _ = self.client.request("POST", "/api/classes", {"name": "WB31"})

        status, template, _ = self.client.request("GET", f"/api/classes/{klass['id']}/template")
        self.assertEqual(status, 200)
        template_book = load_workbook(io.BytesIO(template))
        self.assertEqual([cell.value for cell in template_book.active[1]], ["Vorname", "Klasse"])
        self.assertEqual(template_book.active["B2"].value, "WB31")

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Vorname", "Klasse"])
        sheet.append(["Lena", ""])
        sheet.append(["Yusuf", "WB32"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        status, preview, _ = self.client.request("POST", f"/api/classes/{klass['id']}/import-preview", {
            "file_base64": base64.b64encode(buffer.getvalue()).decode()
        })
        self.assertEqual(status, 200)
        self.assertEqual([row["class_name"] for row in preview["rows"]], ["WB31", "WB32"])
        self.assertFalse(preview["rows"][0]["new_class"])
        self.assertTrue(preview["rows"][1]["new_class"])

        status, result, _ = self.client.request("POST", f"/api/classes/{klass['id']}/import", {
            "rows": preview["rows"]
        })
        self.assertEqual(status, 200)
        self.assertEqual(result["created_classes"], ["WB32"])
        self.assertEqual([row["class_name"] for row in result["created"]], ["WB31", "WB32"])
        wb32 = self.app.db.one("SELECT id FROM classes WHERE name='WB32'")
        self.assertIsNotNone(wb32)
        self.assertEqual(self.app.db.one("SELECT first_name FROM users WHERE class_id=?", (wb32["id"],))["first_name"], "Yusuf")

    def test_global_excel_import_works_without_an_existing_class(self):
        self.setup_teacher()
        status, template, _ = self.client.request("GET", "/api/users/import-template")
        self.assertEqual(status, 200)
        self.assertEqual([cell.value for cell in load_workbook(io.BytesIO(template)).active[1]], ["Vorname", "Klasse"])

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Vorname", "Klasse"])
        sheet.append(["Mia", "GH24"])
        sheet.append(["Noah", "GH25"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        status, preview, _ = self.client.request("POST", "/api/users/import-preview", {
            "file_base64": base64.b64encode(buffer.getvalue()).decode()
        })
        self.assertEqual(status, 200)
        self.assertTrue(all(row["new_class"] for row in preview["rows"]))
        status, result, _ = self.client.request("POST", "/api/users/import", {"rows": preview["rows"]})
        self.assertEqual(status, 200)
        self.assertEqual(result["created_classes"], ["GH24", "GH25"])

        missing_class_book = Workbook()
        missing_class_sheet = missing_class_book.active
        missing_class_sheet.append(["Vorname", "Klasse"])
        missing_class_sheet.append(["Lena", ""])
        missing_class_buffer = io.BytesIO()
        missing_class_book.save(missing_class_buffer)
        status, preview, _ = self.client.request("POST", "/api/users/import-preview", {
            "file_base64": base64.b64encode(missing_class_buffer.getvalue()).decode()
        })
        self.assertEqual(status, 200)
        self.assertIn("Klasse angegeben", preview["rows"][0]["error"])

    def test_class_name_validation_explains_empty_values_and_spaces(self):
        self.setup_teacher()
        status, result, _ = self.client.request("POST", "/api/classes", {"name": "   "})
        self.assertEqual(status, 400)
        self.assertIn("nicht leer", result["error"])
        status, result, _ = self.client.request("POST", "/api/classes", {"name": "Test Klasse"})
        self.assertEqual(status, 400)
        self.assertIn("keine Leerzeichen", result["error"])

    def test_backup_contains_database_uploads_and_reports(self):
        data_dir = Path(self.temp.name) / "backup-source"
        backup_dir = Path(self.temp.name) / "backup-target"
        data_dir.mkdir(); (data_dir / "uploads").mkdir(); (data_dir / "reports").mkdir()
        app = App(Config("127.0.0.1", 8080, data_dir, 25 * 1024 * 1024))
        app.db.execute("INSERT INTO classes(name,created_at) VALUES(?,?)", ("BK32", "2026-07-16T12:00:00+00:00"))
        (data_dir / "uploads" / "test.pdf").write_bytes(b"%PDF-test")
        (data_dir / "reports" / "bericht.pdf").write_bytes(b"%PDF-report")
        with patch.dict(os.environ, {"PK_DATA_DIR": str(data_dir), "PK_BACKUP_DIR": str(backup_dir)}):
            archive = create_backup()
        self.assertTrue(archive.is_file())
        self.assertEqual(archive.stat().st_mode & 0o777, 0o600)
        with tarfile.open(archive, "r:gz") as saved:
            names = set(saved.getnames())
        self.assertIn("projektkontor/projektkontor.sqlite3", names)
        self.assertIn("projektkontor/uploads/test.pdf", names)
        self.assertIn("projektkontor/reports/bericht.pdf", names)

    def test_project_can_be_created_without_start_and_end(self):
        self.setup_teacher()
        _, klass, _ = self.client.request("POST", "/api/classes", {"name": "GH26"})
        _, imported, _ = self.client.request("POST", f"/api/classes/{klass['id']}/import", {"rows": [{"first_name": "Mia"}]})
        lead_id = imported["created"][0]["id"]
        status, project, _ = self.client.request("POST", "/api/projects", {
            "title": "Projekt ohne Zeitraum", "class_id": klass["id"], "description": "Offen",
            "has_start": False, "has_end": False, "start_at": None, "end_at": None,
            "project_lead_id": lead_id, "member_ids": [lead_id],
        })
        self.assertEqual(status, 200)
        status, detail, _ = self.client.request("GET", f"/api/projects/{project['id']}")
        self.assertEqual(status, 200)
        self.assertIsNone(detail["start_at"])
        self.assertIsNone(detail["end_at"])
        self.assertEqual(detail["has_start"], 0)
        self.assertEqual(detail["has_end"], 0)

        status, dated, _ = self.client.request("POST", "/api/projects", {
            "title": "Projekt nur mit Datum", "class_id": klass["id"], "description": "Ohne Uhrzeit",
            "has_start": True, "has_end": True, "start_has_time": False, "end_has_time": False,
            "start_at": "2026-09-01", "end_at": "2026-09-10",
            "project_lead_id": lead_id, "member_ids": [lead_id],
        })
        self.assertEqual(status, 200)
        _, dated_detail, _ = self.client.request("GET", f"/api/projects/{dated['id']}")
        self.assertEqual(dated_detail["start_at"], "2026-09-01")
        self.assertEqual(dated_detail["end_at"], "2026-09-10")
        self.assertEqual(dated_detail["start_has_time"], 0)

        status, timed, _ = self.client.request("POST", "/api/projects", {
            "title": "Projekt mit Uhrzeit", "class_id": klass["id"], "description": "Mit Uhrzeit",
            "has_start": True, "has_end": True, "start_has_time": True, "end_has_time": True,
            "start_at": "2026-09-01T08:15", "end_at": "2026-09-10T13:30",
            "project_lead_id": lead_id, "member_ids": [lead_id],
        })
        self.assertEqual(status, 200)
        _, timed_detail, _ = self.client.request("GET", f"/api/projects/{timed['id']}")
        self.assertEqual(timed_detail["start_at"], "2026-09-01T08:15")
        self.assertEqual(timed_detail["end_at"], "2026-09-10T13:30")
        self.assertEqual(timed_detail["start_has_time"], 1)

    def test_date_only_upload_deadline_defaults_to_end_of_day(self):
        self.setup_teacher()
        _, klass, _ = self.client.request("POST", "/api/classes", {"name": "GH27"})
        _, imported, _ = self.client.request("POST", f"/api/classes/{klass['id']}/import", {"rows": [{"first_name": "Mia"}]})
        lead_id = imported["created"][0]["id"]
        _, project, _ = self.client.request("POST", "/api/projects", {
            "title": "Fristentest", "class_id": klass["id"], "has_start": False, "has_end": False,
            "project_lead_id": lead_id, "member_ids": [lead_id],
        })
        status, task, _ = self.client.request("POST", f"/api/projects/{project['id']}/tasks", {
            "title": "Upload einreichen", "due_at": "2026-10-05"
        })
        self.assertEqual(status, 200)
        _, detail, _ = self.client.request("GET", f"/api/projects/{project['id']}")
        created = next(item for item in detail["tasks"] if item["id"] == task["ids"][0])
        self.assertEqual(created["due_at"], "2026-10-05T23:59")

    def test_teacher_can_update_own_account_securely(self):
        teacher = self.setup_teacher()
        status, result, _ = self.client.request("PATCH", "/api/account", {"first_name": "Herr Müller"})
        self.assertEqual(status, 200)
        self.assertEqual(result["user"]["first_name"], "Herr Müller")
        status, _, _ = self.client.request("PATCH", "/api/account", {"username": "lehrkraft_neu"})
        self.assertEqual(status, 403)
        status, result, _ = self.client.request("PATCH", "/api/account", {
            "username": "lehrkraft_neu", "current_password": TEST_TEACHER_PASSWORD
        })
        self.assertEqual(status, 200)
        self.assertEqual(result["user"]["username"], "lehrkraft_neu")

    def test_admin_manages_paid_licenses_and_license_access(self):
        today = date.today()
        admin = Client(self.app)
        status, _, _ = admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        self.assertEqual(status, 200)
        _, first, _ = admin.request("POST", "/api/teachers", {
            "first_name": "Erika", "username": "lehrkraft_erika"
        })
        _, second, _ = admin.request("POST", "/api/teachers", {
            "first_name": "Markus", "username": "lehrkraft_markus"
        })
        unlicensed = Client(self.app)
        status, denied, _ = unlicensed.request("POST", "/api/login", {
            "login_type": "teacher", "username": "lehrkraft_erika", "password": first["initial_password"]
        })
        self.assertEqual(status, 403)
        self.assertIn("noch nicht freigeschaltet", denied["error"])

        status, rejected, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Schulleitung", "organization": "Beispiel-BK",
            "plan": "single", "seat_limit": 1, "amount_cents": 7900,
            "payment_status": "paid", "status": "active",
            "starts_on": "2026-01-01", "ends_on": "2099-12-31",
            "teacher_ids": [first["id"], second["id"]],
        })
        self.assertEqual(status, 400)
        self.assertIn("höchstens 1", rejected["error"])

        status, rejected, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Beta-Test", "plan": "beta", "seat_limit": 1,
            "amount_cents": 0, "payment_status": "not_required", "status": "active",
            "starts_on": today.isoformat(), "ends_on": (today + timedelta(days=28)).isoformat(),
            "teacher_ids": [first["id"]],
        })
        self.assertEqual(status, 400)
        self.assertIn("gemäß Produktvorgabe", rejected["error"])

        status, license_record, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Frau Beispiel", "organization": "Beispiel-BK",
            "email": "verwaltung@example.org", "billing_address": "Schulweg 1\n12345 Beispielstadt",
            "invoice_reference": "BEST-2026-17", "plan": "department",
            "seat_limit": 5, "amount_cents": 29900, "payment_status": "paid",
            "status": "active", "starts_on": "2026-01-01", "ends_on": "2099-12-31",
            "teacher_ids": [first["id"], second["id"]],
        })
        self.assertEqual(status, 200)
        self.assertEqual(license_record["used_seats"], 2)
        self.assertEqual(license_record["plan"], "department")
        self.assertEqual(self.app.db.one("SELECT license_managed FROM users WHERE id=?", (first["id"],))["license_managed"], 1)

        status, licenses, _ = admin.request("GET", "/api/licenses")
        self.assertEqual(status, 200)
        self.assertEqual(len(licenses), 1)
        self.assertEqual(licenses[0]["organization"], "Beispiel-BK")
        self.assertEqual(licenses[0]["billing_cycle"], "annual")

        status, issuer, _ = admin.request("GET", "/api/invoice-settings")
        self.assertEqual(status, 200)
        self.assertEqual(issuer["business_name"], "PRIMEAdvisory")
        self.assertEqual(issuer["tax_identifier"], "DE453188253")
        status, missing_bank, _ = admin.request("POST", f"/api/licenses/{license_record['id']}/invoice", {})
        self.assertEqual(status, 400)
        self.assertIn("Bankverbindung", missing_bank["error"])
        status, issuer, _ = admin.request("PATCH", "/api/invoice-settings", {
            **issuer, "iban": "DE02120300000000202051", "bic": "BYLADEM1001",
            "bank_name": "Beispielbank",
        })
        self.assertEqual(status, 200)
        self.assertEqual(issuer["tax_mode"], "small_business")
        self.assertEqual(issuer["vat_rate_basis_points"], 0)
        status, invoice, _ = admin.request("POST", f"/api/licenses/{license_record['id']}/invoice", {})
        self.assertEqual(status, 200)
        self.assertRegex(invoice["invoice_number"], r"^PK-\d{4}-0001$")
        stored_invoice = self.app.db.one("SELECT * FROM invoices WHERE id=?", (invoice["id"],))
        self.assertEqual(stored_invoice["net_cents"], stored_invoice["gross_cents"])
        self.assertEqual(stored_invoice["vat_rate_basis_points"], 0)
        self.assertEqual(stored_invoice["vat_cents"], 0)
        self.assertEqual(stored_invoice["tax_note"], "Gemäß § 19 UStG wird keine Umsatzsteuer berechnet.")
        status, invoice_pdf, headers = admin.request("GET", f"/api/invoices/{invoice['id']}")
        self.assertEqual(status, 200)
        self.assertEqual(headers["Content-Type"], "application/pdf")
        self.assertTrue(invoice_pdf.startswith(b"%PDF-"))
        self.assertIn(b"%%EOF", invoice_pdf[-1024:])
        self.assertGreater(len(invoice_pdf), 2000)
        object.__setattr__(self.app.config, "smtp_host", "mail.gmx.net")
        object.__setattr__(self.app.config, "smtp_username", "konto@gmx.de")
        object.__setattr__(self.app.config, "smtp_password", "anwendungspasswort")
        object.__setattr__(self.app.config, "smtp_sender", "abweichend@example.org")
        with patch("projektkontor.server.smtplib.SMTP_SSL") as smtp:
            smtp.return_value.__enter__.return_value = smtp.return_value
            status, sent, _ = admin.request("POST", f"/api/invoices/{invoice['id']}/email", {})
            self.assertEqual(status, 200)
            self.assertTrue(sent["ok"])
            message = smtp.return_value.send_message.call_args.args[0]
        self.assertIn("konto@gmx.de", message["From"])
        self.assertEqual(message["To"], "verwaltung@example.org")
        self.assertEqual(message["Reply-To"], "abweichend@example.org")
        self.assertEqual(message.get_content_maintype(), "multipart")

        teacher = Client(self.app)
        status, _, _ = teacher.request("POST", "/api/login", {
            "login_type": "teacher", "username": "lehrkraft_erika", "password": first["initial_password"]
        })
        self.assertEqual(status, 200)
        status, bootstrap, _ = teacher.request("GET", "/api/bootstrap")
        self.assertEqual(status, 200)
        self.assertEqual(bootstrap["user"]["first_name"], "Erika")

        status, updated, _ = admin.request("PATCH", f"/api/licenses/{license_record['id']}", {
            "status": "suspended",
        })
        self.assertEqual(status, 200)
        self.assertEqual(updated["status"], "suspended")
        status, bootstrap, _ = teacher.request("GET", "/api/bootstrap")
        self.assertEqual(status, 200)
        self.assertIsNone(bootstrap["user"])
        status, denied, _ = teacher.request("POST", "/api/login", {
            "login_type": "teacher", "username": "lehrkraft_erika", "password": TEST_TEACHER_PASSWORD
        })
        self.assertEqual(status, 403)
        self.assertIn("noch nicht freigeschaltet", denied["error"])
        self.assertEqual(self.app.db.one("SELECT active FROM users WHERE id=?", (first["id"],))["active"], 1)

        status, _, _ = admin.request("PATCH", f"/api/licenses/{license_record['id']}", {
            "status": "active",
        })
        self.assertEqual(status, 200)
        status, _, _ = teacher.request("POST", "/api/login", {
            "login_type": "teacher", "username": "lehrkraft_erika", "password": TEST_TEACHER_PASSWORD
        })
        self.assertEqual(status, 200)
        self.app.db.execute(
            "UPDATE licenses SET starts_on=?,ends_on=?,status='active' WHERE id=?",
            ("2020-01-01", (today - timedelta(days=1)).isoformat(), license_record["id"]),
        )
        status, bootstrap, _ = teacher.request("GET", "/api/bootstrap")
        self.assertEqual(status, 200)
        self.assertIsNone(bootstrap["user"])
        self.assertEqual(
            self.app.db.one("SELECT status FROM licenses WHERE id=?", (license_record["id"],))["status"],
            "expired",
        )

    def test_teacher_creation_can_create_or_select_a_standard_license(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        status, monthly_teacher, _ = admin.request("POST", "/api/teachers", {
            "first_name": "Frau Monat", "username": "lehrkraft_monat",
            "license_selection": "new:monthly",
            "license_customer_name": "Frau Monat",
            "license_organization": "Monat-BK",
            "license_email": "monat@example.org",
            "license_billing_address": "Monatstraße 1\n45127 Essen",
        })
        self.assertEqual(status, 200)
        monthly_license = self.app.db.one(
            """SELECT l.id,l.status,l.starts_on,l.ends_on,o.plan,o.billing_cycle,o.amount_cents
               FROM license_teachers lt JOIN licenses l ON l.id=lt.license_id
               JOIN license_orders o ON o.id=l.order_id WHERE lt.teacher_id=?""",
            (monthly_teacher["id"],),
        )
        self.assertEqual(monthly_license["plan"], "single")
        self.assertEqual(monthly_license["billing_cycle"], "monthly")
        self.assertEqual(monthly_license["amount_cents"], 890)
        self.assertEqual(monthly_license["status"], "active")

        status, unassigned, _ = admin.request("POST", "/api/teachers", {
            "first_name": "Herr Frei", "username": "lehrkraft_frei"
        })
        self.assertEqual(status, 200)
        status, _, _ = admin.request("PATCH", f"/api/teachers/{unassigned['id']}", {
            "first_name": "Herr Frei", "username": "lehrkraft_frei", "active": True,
            "license_selection": f"existing:{monthly_license['id']}",
        })
        self.assertEqual(status, 409)
        department_status, department, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Fachbereich", "organization": "Monat-BK",
            "plan": "department", "seat_limit": 5, "amount_cents": 29900,
            "payment_status": "open", "status": "active",
            "starts_on": "2020-01-01", "ends_on": "2099-12-31", "teacher_ids": [],
        })
        self.assertEqual(department_status, 200)
        status, _, _ = admin.request("PATCH", f"/api/teachers/{unassigned['id']}", {
            "first_name": "Herr Frei", "username": "lehrkraft_frei", "active": True,
            "license_selection": f"existing:{department['id']}",
        })
        self.assertEqual(status, 200)
        self.assertTrue(self.app.teacher_license_valid(unassigned["id"]))

    def test_license_creation_can_atomically_create_and_assign_teacher(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        status, license_record, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Frau Direkt", "organization": "Direkt-BK",
            "plan": "single", "billing_cycle": "annual", "seat_limit": 1,
            "amount_cents": 7900, "payment_status": "open", "status": "active",
            "starts_on": "2020-01-01", "ends_on": "2099-12-31",
            "teacher_ids": [], "create_teacher": True,
            "new_teacher_first_name": "Frau Direkt",
            "new_teacher_username": "lehrkraft_direkt",
            "new_teacher_email": "direkt@example.org",
        })
        self.assertEqual(status, 200)
        created = license_record["created_teacher"]
        self.assertEqual(created["username"], "lehrkraft_direkt")
        self.assertTrue(created["initial_password"])
        self.assertEqual(license_record["used_seats"], 1)
        self.assertEqual(license_record["teachers"][0]["id"], created["id"])
        self.assertTrue(self.app.teacher_license_valid(created["id"]))
        object.__setattr__(self.app.config, "smtp_host", "mail.gmx.net")
        object.__setattr__(self.app.config, "smtp_username", "konto@gmx.de")
        object.__setattr__(self.app.config, "smtp_password", "anwendungspasswort")
        with patch("projektkontor.server.smtplib.SMTP_SSL") as smtp:
            smtp.return_value.__enter__.return_value = smtp.return_value
            mail_status, mailed, _ = admin.request(
                "POST", f"/api/teachers/{created['id']}/initial-credentials/email",
                {"initial_password": created["initial_password"]},
            )
            message = smtp.return_value.send_message.call_args.args[0]
        self.assertEqual(mail_status, 200)
        self.assertTrue(mailed["ok"])
        self.assertEqual(message["To"], "direkt@example.org")
        self.assertIn("lehrkraft_direkt", message.get_content())
        self.assertIn(created["initial_password"], message.get_content())
        self.assertTrue(
            self.app.db.one("SELECT initial_credentials_emailed_at FROM users WHERE id=?", (created["id"],))[
                "initial_credentials_emailed_at"
            ]
        )

        duplicate_status, duplicate, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Doppelt", "plan": "single", "billing_cycle": "annual",
            "seat_limit": 1, "amount_cents": 7900, "payment_status": "open",
            "status": "active", "starts_on": "2020-01-01", "ends_on": "2099-12-31",
            "teacher_ids": [], "create_teacher": True,
            "new_teacher_first_name": "Noch einmal",
            "new_teacher_username": "lehrkraft_direkt",
            "new_teacher_email": "doppelt@example.org",
        })
        self.assertEqual(duplicate_status, 409)
        self.assertIn("bereits vergeben", duplicate["error"])
        self.assertEqual(self.app.db.one("SELECT COUNT(*) count FROM licenses")["count"], 1)
        self.assertEqual(self.app.db.one("SELECT COUNT(*) count FROM license_orders")["count"], 1)

    def test_license_can_be_archived_or_permanently_deleted_with_invoices(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        _, settings, _ = admin.request("GET", "/api/invoice-settings")
        admin.request("PATCH", "/api/invoice-settings", {
            **settings, "iban": "DE02120300000000202051", "bic": "BYLADEM1001",
            "bank_name": "Beispielbank",
        })
        status, license_record, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Archiv Test", "organization": "Archiv-BK",
            "email": "archiv@example.org", "billing_address": "Testweg 1\n45127 Essen",
            "plan": "single", "billing_cycle": "annual", "seat_limit": 1,
            "amount_cents": 7900, "payment_status": "open", "status": "active",
            "starts_on": "2020-01-01", "ends_on": "2099-12-31", "teacher_ids": [],
        })
        self.assertEqual(status, 200)
        invoice_status, invoice, _ = admin.request(
            "POST", f"/api/licenses/{license_record['id']}/invoice", {}
        )
        self.assertEqual(invoice_status, 200)
        stored_name = self.app.db.one("SELECT stored_name FROM invoices WHERE id=?", (invoice["id"],))["stored_name"]
        invoice_path = self.app.config.report_dir / stored_name
        self.assertTrue(invoice_path.is_file())

        archive_status, archived, _ = admin.request(
            "POST", f"/api/licenses/{license_record['id']}/archive", {}
        )
        self.assertEqual(archive_status, 200)
        self.assertEqual(archived["status"], "cancelled")
        self.assertTrue(archived["archived_at"])
        self.assertTrue(invoice_path.is_file())

        rejected_status, rejected, _ = admin.request(
            "DELETE", f"/api/licenses/{license_record['id']}", {}
        )
        self.assertEqual(rejected_status, 400)
        self.assertIn("endgültige Löschen", rejected["error"])
        delete_status, deleted, _ = admin.request(
            "DELETE", f"/api/licenses/{license_record['id']}", {"confirm_permanent_delete": True}
        )
        self.assertEqual(delete_status, 200)
        self.assertTrue(deleted["ok"])
        self.assertIsNone(self.app.db.one("SELECT id FROM licenses WHERE id=?", (license_record["id"],)))
        self.assertIsNone(self.app.db.one("SELECT id FROM invoices WHERE id=?", (invoice["id"],)))
        retained = self.app.db.one(
            "SELECT invoice_id,retired_at FROM invoice_number_registry WHERE invoice_number=?",
            (invoice["invoice_number"],),
        )
        self.assertIsNone(retained["invoice_id"])
        self.assertTrue(retained["retired_at"])
        self.assertFalse(invoice_path.exists())

    def test_teacher_accounts_can_be_archived_restored_or_safely_deleted(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        _, unused, _ = admin.request("POST", "/api/teachers", {
            "first_name": "Versehentlich", "username": "lehrkraft_versehen"
        })
        status, _, _ = admin.request("DELETE", f"/api/teachers/{unused['id']}", {})
        self.assertEqual(status, 400)
        status, deleted, _ = admin.request(
            "DELETE", f"/api/teachers/{unused['id']}", {"confirm_permanent_delete": True}
        )
        self.assertEqual(status, 200)
        self.assertEqual(deleted["deleted_id"], unused["id"])
        self.assertIsNone(self.app.db.one("SELECT id FROM users WHERE id=?", (unused["id"],)))

        _, used, _ = admin.request("POST", "/api/teachers", {
            "first_name": "Genutzt", "username": "lehrkraft_genutzt"
        })
        self.app.db.execute(
            "UPDATE users SET last_login_at=? WHERE id=?", ("2026-07-28T07:00:00+00:00", used["id"])
        )
        status, blocked, _ = admin.request(
            "DELETE", f"/api/teachers/{used['id']}", {"confirm_permanent_delete": True}
        )
        self.assertEqual(status, 409)
        self.assertIn("nur archiviert", blocked["error"])
        status, archived, _ = admin.request("POST", f"/api/teachers/{used['id']}/archive", {})
        self.assertEqual(status, 200)
        self.assertTrue(archived["archived_at"])
        account = self.app.db.one("SELECT active,archived_at FROM users WHERE id=?", (used["id"],))
        self.assertEqual(account["active"], 0)
        self.assertTrue(account["archived_at"])
        status, restored, _ = admin.request("POST", f"/api/teachers/{used['id']}/restore", {})
        self.assertEqual(status, 200)
        self.assertTrue(restored["restored"])
        account = self.app.db.one("SELECT active,archived_at FROM users WHERE id=?", (used["id"],))
        self.assertEqual(account, {"active": 1, "archived_at": None})

    def test_invoice_archive_and_delete_never_reuses_invoice_numbers(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        _, license_record, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Testkundin", "organization": "Test-BK",
            "billing_address": "Testweg 1\n45127 Essen", "plan": "beta",
            "amount_cents": 0, "seat_limit": 1, "payment_status": "not_required",
            "status": "active", "starts_on": "2026-07-01", "ends_on": "2026-07-14",
            "teacher_ids": [],
        })
        _, first, _ = admin.request(
            "POST", f"/api/licenses/{license_record['id']}/invoice", {"confirm_zero_invoice": True}
        )
        _, second, _ = admin.request(
            "POST", f"/api/licenses/{license_record['id']}/invoice", {"confirm_zero_invoice": True}
        )
        status, archived, _ = admin.request("POST", f"/api/invoices/{first['id']}/archive", {})
        self.assertEqual(status, 200)
        self.assertEqual(archived["status"], "cancelled")
        self.assertTrue(archived["archived_at"])
        status, denied, _ = admin.request("POST", f"/api/invoices/{first['id']}/email", {})
        self.assertEqual(status, 409)
        self.assertIn("stornierte", denied["error"])

        status, deleted, _ = admin.request(
            "DELETE", f"/api/invoices/{second['id']}", {"confirm_permanent_delete": True}
        )
        self.assertEqual(status, 200)
        self.assertEqual(deleted["retained_invoice_number"], second["invoice_number"])
        registry = self.app.db.one(
            "SELECT invoice_id,retired_at FROM invoice_number_registry WHERE invoice_number=?",
            (second["invoice_number"],),
        )
        self.assertIsNone(registry["invoice_id"])
        self.assertTrue(registry["retired_at"])
        _, third, _ = admin.request(
            "POST", f"/api/licenses/{license_record['id']}/invoice", {"confirm_zero_invoice": True}
        )
        self.assertNotEqual(third["invoice_number"], second["invoice_number"])
        self.assertTrue(third["invoice_number"].endswith("0003"))

    def test_beta_invoice_conversion_history_and_payment_terms(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        status, settings, _ = admin.request("GET", "/api/invoice-settings")
        self.assertEqual(status, 200)
        self.assertEqual(settings["payment_terms_days"], 0)
        status, invalid, _ = admin.request("PATCH", "/api/invoice-settings", {
            **settings, "payment_terms_days": 21,
        })
        self.assertEqual(status, 400)
        self.assertIn("Sofort, 7, 14 oder 30 Tage", invalid["error"])

        status, beta, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Frau Test", "organization": "Beta-BK",
            "email": "beta@example.org", "billing_address": "Testweg 1\n45127 Essen",
            "plan": "beta", "amount_cents": 0, "seat_limit": 1,
            "payment_status": "not_required", "status": "active",
            "starts_on": "2026-07-01", "ends_on": "2026-07-14", "teacher_ids": [],
        })
        self.assertEqual(status, 200)
        status, warning, _ = admin.request("POST", f"/api/licenses/{beta['id']}/invoice", {})
        self.assertEqual(status, 400)
        self.assertIn("Nullrechnung", warning["error"])
        status, zero_invoice, _ = admin.request("POST", f"/api/licenses/{beta['id']}/invoice", {
            "confirm_zero_invoice": True,
        })
        self.assertEqual(status, 200)
        self.assertEqual(zero_invoice["gross_cents"], 0)
        self.assertEqual(zero_invoice["issued_on"], zero_invoice["due_on"])
        self.assertEqual(self.app.license_record(beta["id"])["payment_status"], "not_required")

        status, paid, _ = admin.request("PATCH", f"/api/licenses/{beta['id']}", {
            "plan": "single", "billing_cycle": "annual", "amount_cents": 7900,
            "payment_status": "open", "status": "active",
            "starts_on": "2026-07-29", "ends_on": "2027-07-28",
        })
        self.assertEqual(status, 200)
        self.assertEqual(paid["plan"], "single")
        self.assertEqual(paid["amount_cents"], 7900)
        self.assertEqual(paid["history"][0]["event_type"], "converted")
        self.assertEqual(paid["history"][0]["from_plan"], "beta")
        self.assertEqual(paid["history"][0]["to_plan"], "single")
        self.assertTrue(paid["history"][0]["changed_at"])

        status, settings, _ = admin.request("PATCH", "/api/invoice-settings", {
            **settings, "payment_terms_days": 7, "iban": "DE02120300000000202051",
            "bic": "BYLADEM1001", "bank_name": "Beispielbank",
        })
        self.assertEqual(status, 200)
        status, paid_invoice, _ = admin.request("POST", f"/api/licenses/{beta['id']}/invoice", {})
        self.assertEqual(status, 200)
        self.assertEqual(paid_invoice["gross_cents"], 7900)
        self.assertEqual(
            (date.fromisoformat(paid_invoice["due_on"]) - date.fromisoformat(paid_invoice["issued_on"])).days,
            7,
        )
        status, licenses, _ = admin.request("GET", "/api/licenses")
        self.assertEqual(status, 200)
        updated = next(item for item in licenses if item["id"] == beta["id"])
        self.assertEqual(len(updated["invoices"]), 2)

    def test_paid_active_license_is_locked_and_can_schedule_follow_up(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        _, teacher, _ = admin.request("POST", "/api/teachers", {
            "first_name": "Erika", "username": "lehrkraft_erika"
        })
        today = date.today()
        source_end = today + timedelta(days=30)
        status, source, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Frau Vertrag", "organization": "Vertrags-BK",
            "email": "vertrag@example.org", "billing_address": "Testweg 1\n45127 Essen",
            "plan": "single", "billing_cycle": "annual", "amount_cents": 7900,
            "seat_limit": 1, "payment_status": "paid", "status": "active",
            "starts_on": today.isoformat(), "ends_on": source_end.isoformat(),
            "teacher_ids": [teacher["id"]],
        })
        self.assertEqual(status, 200)
        status, denied, _ = admin.request("PATCH", f"/api/licenses/{source['id']}", {
            "plan": "school", "amount_cents": 59900, "seat_limit": 15,
        })
        self.assertEqual(status, 409)
        self.assertIn("Folgelizenz", denied["error"])

        follow_start = source_end + timedelta(days=1)
        follow_end = follow_start + timedelta(days=364)
        status, follow_up, _ = admin.request("POST", "/api/licenses", {
            "follow_up_of": source["id"], "customer_name": "Frau Vertrag",
            "organization": "Vertrags-BK", "email": "vertrag@example.org",
            "billing_address": "Testweg 1\n45127 Essen", "plan": "school",
            "billing_cycle": "annual", "amount_cents": 59900, "seat_limit": 15,
            "payment_status": "paid", "status": "active",
            "starts_on": follow_start.isoformat(), "ends_on": follow_end.isoformat(),
            "teacher_ids": [teacher["id"]],
        })
        self.assertEqual(status, 200)
        self.assertEqual(follow_up["status"], "draft")
        self.assertEqual(follow_up["payment_status"], "open")
        self.assertEqual(follow_up["follow_up_of"], source["id"])
        self.assertEqual(self.app.license_record(source["id"])["status"], "active")
        status, follow_up, _ = admin.request("PATCH", f"/api/licenses/{follow_up['id']}", {
            "amount_cents": 64900,
        })
        self.assertEqual(status, 200)
        self.assertEqual(follow_up["amount_cents"], 64900)
        assignments = self.app.db.all(
            "SELECT license_id FROM license_teachers WHERE teacher_id=? ORDER BY license_id",
            (teacher["id"],),
        )
        self.assertEqual([row["license_id"] for row in assignments], [source["id"], follow_up["id"]])
        self.app.db.execute(
            "UPDATE licenses SET ends_on=? WHERE id=?", ((today - timedelta(days=1)).isoformat(), source["id"])
        )
        self.app.db.execute(
            "UPDATE licenses SET starts_on=?,ends_on=? WHERE id=?",
            (today.isoformat(), (today + timedelta(days=364)).isoformat(), follow_up["id"]),
        )
        self.app.refresh_expired_licenses()
        self.assertEqual(self.app.license_record(source["id"])["status"], "expired")
        self.assertEqual(self.app.license_record(follow_up["id"])["status"], "active")
        self.assertTrue(self.app.teacher_license_valid(teacher["id"]))

    def test_follow_up_reminder_activation_and_invoice_are_automatic_and_idempotent(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        settings = self.app.db.one("SELECT * FROM invoice_settings WHERE id=1")
        admin.request("PATCH", "/api/invoice-settings", {
            **settings, "iban": "DE02120300000000202051",
            "bic": "BYLADEM1001", "bank_name": "Beispielbank",
        })
        today = date.today()
        _, source, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Automatik", "organization": "Automatik-BK",
            "email": "rechnung@example.org", "billing_address": "Testweg 1\n45127 Essen",
            "plan": "single", "billing_cycle": "annual", "amount_cents": 7900,
            "seat_limit": 1, "payment_status": "paid", "status": "active",
            "starts_on": (today - timedelta(days=360)).isoformat(),
            "ends_on": (today + timedelta(days=5)).isoformat(), "teacher_ids": [],
        })
        _, follow_up, _ = admin.request("POST", "/api/licenses", {
            "follow_up_of": source["id"], "customer_name": "Automatik",
            "organization": "Automatik-BK", "email": "rechnung@example.org",
            "billing_address": "Testweg 1\n45127 Essen", "plan": "single",
            "billing_cycle": "annual", "amount_cents": 8900, "seat_limit": 1,
            "starts_on": (today + timedelta(days=6)).isoformat(),
            "ends_on": (today + timedelta(days=370)).isoformat(), "teacher_ids": [],
        })
        result = self.app.run_license_automation()
        self.assertEqual(result["reminded"], 1)
        self.assertTrue(self.app.license_record(follow_up["id"])["follow_up_reminded_at"])

        self.app.db.execute(
            "UPDATE licenses SET ends_on=? WHERE id=?",
            ((today - timedelta(days=1)).isoformat(), source["id"]),
        )
        self.app.db.execute(
            "UPDATE licenses SET starts_on=? WHERE id=?", (today.isoformat(), follow_up["id"])
        )
        first = self.app.run_license_automation()
        second = self.app.run_license_automation()
        activated = self.app.license_record(follow_up["id"])
        self.assertEqual(first["created"], 1)
        self.assertEqual(second["created"], 0)
        self.assertEqual(activated["status"], "active")
        self.assertEqual(len(activated["invoices"]), 1)
        self.assertTrue(activated["invoice_automation_completed_at"])

    def test_invoice_paid_action_syncs_license_payment_status(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        settings = self.app.db.one("SELECT * FROM invoice_settings WHERE id=1")
        admin.request("PATCH", "/api/invoice-settings", {
            **settings, "iban": "DE02120300000000202051",
        })
        _, license_record, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Zahlstatus", "billing_address": "Testweg 1\n45127 Essen",
            "plan": "single", "billing_cycle": "annual", "amount_cents": 7900,
            "seat_limit": 1, "payment_status": "open", "status": "active",
            "starts_on": date.today().isoformat(),
            "ends_on": (date.today() + timedelta(days=364)).isoformat(), "teacher_ids": [],
        })
        status, invoice, _ = admin.request(
            "POST", f"/api/licenses/{license_record['id']}/invoice", {}
        )
        self.assertEqual(status, 200)
        status, paid, _ = admin.request("POST", f"/api/invoices/{invoice['id']}/paid", {})
        self.assertEqual(status, 200)
        self.assertEqual(paid["status"], "paid")
        self.assertEqual(self.app.license_record(license_record["id"])["payment_status"], "paid")

    def test_source_license_cannot_be_deleted_before_its_follow_up(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        today = date.today()
        _, source, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Vormerkung", "plan": "single", "billing_cycle": "annual",
            "amount_cents": 7900, "seat_limit": 1, "payment_status": "paid", "status": "active",
            "starts_on": today.isoformat(), "ends_on": (today + timedelta(days=30)).isoformat(),
            "teacher_ids": [],
        })
        _, follow_up, _ = admin.request("POST", "/api/licenses", {
            "follow_up_of": source["id"], "customer_name": "Vormerkung", "plan": "school",
            "billing_cycle": "annual", "amount_cents": 59900, "seat_limit": 15,
            "starts_on": (today + timedelta(days=31)).isoformat(),
            "ends_on": (today + timedelta(days=395)).isoformat(), "teacher_ids": [],
        })
        status, error, _ = admin.request(
            "DELETE", f"/api/licenses/{source['id']}", {"confirm_permanent_delete": True}
        )
        self.assertEqual(status, 409)
        self.assertIn("Vormerkung", error["error"])
        status, _, _ = admin.request(
            "DELETE", f"/api/licenses/{follow_up['id']}", {"confirm_permanent_delete": True}
        )
        self.assertEqual(status, 200)
        status, _, _ = admin.request(
            "DELETE", f"/api/licenses/{source['id']}", {"confirm_permanent_delete": True}
        )
        self.assertEqual(status, 200)

    def test_monthly_courtesy_cancellation_fully_cancels_invoice(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        settings = self.app.db.one("SELECT * FROM invoice_settings WHERE id=1")
        admin.request("PATCH", "/api/invoice-settings", {**settings, "iban": "DE02120300000000202051"})
        today = date.today()
        _, license_record, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Monatskulanz", "billing_address": "Testweg 1\n45127 Essen",
            "plan": "single", "billing_cycle": "monthly", "amount_cents": 890,
            "seat_limit": 1, "payment_status": "open", "status": "active",
            "starts_on": today.isoformat(),
            "ends_on": (self.app.add_months(today, 1) - timedelta(days=1)).isoformat(),
            "teacher_ids": [],
        })
        _, invoice, _ = admin.request("POST", f"/api/licenses/{license_record['id']}/invoice", {})
        admin.request("POST", f"/api/invoices/{invoice['id']}/paid", {})
        status, cancelled, _ = admin.request(
            "POST", f"/api/licenses/{license_record['id']}/courtesy-cancel", {
                "effective_on": today.isoformat(), "notes": "Kulanzfall",
                "confirm_courtesy_cancellation": True,
            },
        )
        self.assertEqual(status, 200)
        self.assertEqual(cancelled["status"], "cancelled")
        self.assertEqual(cancelled["amount_cents"], 0)
        self.assertEqual(cancelled["payment_status"], "refunded")
        self.assertEqual(cancelled["refund_due_cents"], 890)
        self.assertEqual(cancelled["cancellation"]["mode"], "monthly_full")
        self.assertTrue(cancelled["invoices"][0]["archived_at"])

    def test_annual_courtesy_cancellation_keeps_only_completed_months(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        settings = self.app.db.one("SELECT * FROM invoice_settings WHERE id=1")
        admin.request("PATCH", "/api/invoice-settings", {**settings, "iban": "DE02120300000000202051"})
        today = date.today()
        start = self.app.add_months(today, -6)
        end = self.app.add_months(start, 12) - timedelta(days=1)
        _, license_record, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Jahreskulanz", "billing_address": "Testweg 1\n45127 Essen",
            "plan": "single", "billing_cycle": "annual", "amount_cents": 12000,
            "seat_limit": 1, "payment_status": "open", "status": "active",
            "starts_on": start.isoformat(), "ends_on": end.isoformat(), "teacher_ids": [],
        })
        _, original, _ = admin.request("POST", f"/api/licenses/{license_record['id']}/invoice", {})
        admin.request("POST", f"/api/invoices/{original['id']}/paid", {})
        status, cancelled, _ = admin.request(
            "POST", f"/api/licenses/{license_record['id']}/courtesy-cancel", {
                "effective_on": today.isoformat(), "confirm_courtesy_cancellation": True,
            },
        )
        self.assertEqual(status, 200)
        self.assertEqual(cancelled["status"], "cancelled")
        self.assertEqual(cancelled["amount_cents"], 6000)
        self.assertEqual(cancelled["cancellation"]["retained_amount_cents"], 6000)
        self.assertEqual(cancelled["refund_due_cents"], 6000)
        active_invoices = [row for row in cancelled["invoices"] if not row["archived_at"]]
        self.assertEqual(len(active_invoices), 1)
        self.assertEqual(active_invoices[0]["gross_cents"], 6000)
        self.assertEqual(active_invoices[0]["status"], "paid")

    def test_assigning_a_teacher_to_another_license_transfers_the_assignment(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        _, teacher, _ = admin.request("POST", "/api/teachers", {
            "first_name": "Erika", "username": "lehrkraft_erika"
        })
        common = {
            "customer_name": "Erika", "plan": "single", "billing_cycle": "annual",
            "seat_limit": 1, "amount_cents": 7900, "payment_status": "paid",
            "status": "active", "starts_on": "2026-01-01", "ends_on": "2099-12-31",
        }
        status, first, _ = admin.request("POST", "/api/licenses", {
            **common, "organization": "Alte Schule", "teacher_ids": [teacher["id"]],
        })
        self.assertEqual(status, 200)
        status, second, _ = admin.request("POST", "/api/licenses", {
            **common, "organization": "Neue Schule", "teacher_ids": [],
        })
        self.assertEqual(status, 200)
        status, second, _ = admin.request("PATCH", f"/api/licenses/{second['id']}", {
            "organization": "Neue Schule – aktualisiert", "teacher_ids": [teacher["id"]],
        })
        self.assertEqual(status, 200)
        self.assertEqual(second["organization"], "Neue Schule – aktualisiert")
        self.assertEqual(second["used_seats"], 1)
        refreshed_first = self.app.license_record(first["id"])
        self.assertEqual(refreshed_first["used_seats"], 0)
        assignment = self.app.db.one(
            "SELECT license_id FROM license_teachers WHERE teacher_id=?", (teacher["id"],)
        )
        self.assertEqual(assignment["license_id"], second["id"])

    def test_license_product_catalog_updates_only_future_licenses(self):
        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        status, products, _ = admin.request("GET", "/api/license-products")
        self.assertEqual(status, 200)
        monthly = next(product for product in products if product["product_key"] == "single_monthly")
        self.assertEqual(monthly["amount_cents"], 890)
        _, first_teacher, _ = admin.request("POST", "/api/teachers", {
            "first_name": "Altpreis", "username": "lehrkraft_altpreis",
            "license_selection": "new-product:single_monthly",
        })
        first_license = self.app.db.one(
            """SELECT o.amount_cents,l.starts_on,l.ends_on FROM license_teachers lt
                 JOIN licenses l ON l.id=lt.license_id JOIN license_orders o ON o.id=l.order_id
                WHERE lt.teacher_id=?""", (first_teacher["id"],),
        )
        self.assertEqual(first_license["amount_cents"], 890)

        status, updated, _ = admin.request("PATCH", "/api/license-products/single_monthly", {
            "name": "Einzellizenz Flex", "amount_cents": 990, "seat_limit": 1,
            "duration_value": 2, "duration_unit": "months", "active": True,
        })
        self.assertEqual(status, 200)
        self.assertEqual(updated["amount_cents"], 990)
        _, second_teacher, _ = admin.request("POST", "/api/teachers", {
            "first_name": "Neupreis", "username": "lehrkraft_neupreis",
            "license_selection": "new-product:single_monthly",
        })
        second_license = self.app.db.one(
            """SELECT o.amount_cents,l.starts_on,l.ends_on FROM license_teachers lt
                 JOIN licenses l ON l.id=lt.license_id JOIN license_orders o ON o.id=l.order_id
                WHERE lt.teacher_id=?""", (second_teacher["id"],),
        )
        self.assertEqual(second_license["amount_cents"], 990)
        self.assertEqual(first_license["amount_cents"], 890)
        self.assertGreater(
            (date.fromisoformat(second_license["ends_on"]) - date.fromisoformat(second_license["starts_on"])).days,
            45,
        )

    def test_security_headers_and_malformed_image_rejection(self):
        status, _, headers = self.client.request("GET", "/api/health")
        self.assertEqual(status, 200)
        self.assertEqual(headers["X-Content-Type-Options"], "nosniff")
        self.assertEqual(headers["X-Frame-Options"], "DENY")
        self.assertIn("object-src 'none'", headers["Content-Security-Policy"])
        self.assertIn("noindex", headers["X-Robots-Tag"])
        status, app_script, _ = self.client.request("GET", "/app.js")
        self.assertEqual(status, 200)
        self.assertNotIn(b"onclick=", app_script)
        self.assertIn(b'href="#/admin/users"', app_script)
        self.assertIn(b'href="#/projects/${id}/${key}"', app_script)
        self.assertIn(b'id="add-user-from-overview"', app_script)
        self.assertIn(b"chooseManualAccountClassDialog", app_script)
        self.assertIn(b'href="#/admin/licenses">Bestellungen &amp; Lizenzen</a>', app_script)
        self.assertIn(b"async function renderLicenses(options={})", app_script)
        self.assertIn(b"Zahlungsziel verpasst", app_script)
        self.assertIn(b"paymentDashboardMarkup", app_script)
        self.assertIn(b"occupiedTeachersMarkup", app_script)
        self.assertIn(b'id="license-workspace"', app_script)
        self.assertIn(b"Aktuell einer aktiven Lizenz zugeordnete Lehrkraftzug\xc3\xa4nge", app_script)
        self.assertIn(b"E-Mail schreiben", app_script)
        self.assertIn(b"openLicenseId:license.id", app_script)
        self.assertIn(b"Folgelizenz planen", app_script)
        self.assertIn(b"Aktive Bezahl-Lizenz gesch\xc3\xbctzt", app_script)
        self.assertIn(b"Beta-Konvertierung m\xc3\xb6glich", app_script)
        self.assertIn(b"Vorgemerkt", app_script)
        self.assertIn(b"Inaktive Lizenzen", app_script)
        self.assertIn(b"Rechnung ausstehend", app_script)
        self.assertIn(b"Zugeordnete Lehrkraftzug\xc3\xa4nge", app_script)
        self.assertIn(b"Rechnungssteller", app_script)
        self.assertIn(b"new-product:", app_script)
        self.assertIn(b"/api/invoice-settings", app_script)
        self.assertIn(b"Rechnung erstellen", app_script)
        self.assertIn(b'class="button small download-invoice"', app_script)
        self.assertIn(b"Lizenzdetails", app_script)
        self.assertIn(b"teacher-card-list", app_script)
        self.assertIn(b'<details class="card teacher-card', app_script)
        self.assertIn(b"Zugang bearbeiten", app_script)
        self.assertIn(b"billing-address-grid", app_script)
        self.assertIn(b"Stra\xc3\x9fe und Hausnummer", app_script)
        self.assertIn(b"teacher-search", app_script)
        self.assertIn(b"Lizenz direkt zuordnen", app_script)
        self.assertIn(b"teacher-archive-section", app_script)
        self.assertIn(b"Archivieren oder l\xc3\xb6schen", app_script)
        self.assertIn(b'<option value="0"', app_script)
        self.assertIn(b">Sofort</option>", app_script)
        self.assertIn(b">30 Tage</option>", app_script)
        self.assertIn(b"strukturierte XRechnung", app_script)
        self.assertIn(b'<button class="button primary" id="add-user">Zugang manuell anlegen</button>', app_script)
        with self.assertRaises(Exception) as rejected:
            self.app.validated_upload({
                "name": "scheinbild.png", "media_type": "image/png",
                "file_base64": base64.b64encode(b"\x89PNG\r\n\x1a\nkein-echtes-bild").decode(),
            })
        self.assertIn("beschädigt", str(rejected.exception))
        status, robots, headers = self.client.request("GET", "/robots.txt")
        self.assertEqual(status, 200)
        self.assertIn(b"Allow: /", robots)
        self.assertIn(b"Disallow: /api/", robots)
        self.assertNotIn("X-Robots-Tag", headers)

    def test_public_landing_page_and_protected_login_are_separated(self):
        status, landing, headers = self.client.request("GET", "/")
        self.assertEqual(status, 200)
        self.assertIn(b"Berufliche Handlungssituationen vollst", landing)
        self.assertIn(b"webbasierte Arbeitsumgebung", landing)
        self.assertIn(b"Kanban-Board", landing)
        self.assertIn(b'id="contact-form"', landing)
        self.assertNotIn(b"mailto:", landing)
        self.assertNotIn(b"gmx.de", landing)
        self.assertIn(b"Praxiserfahrung", landing)
        self.assertIn(b"ERP-Umfeld", landing)
        self.assertIn(b"Pr\xc3\xbcfstatus", landing)
        self.assertIn(b"Der Projektbericht", landing)
        self.assertIn(b"Produkteinblick", landing)
        self.assertIn(b'href="#vorteile">Vorteile</a>', landing)
        self.assertIn(b"Der entscheidende Unterschied", landing)
        self.assertIn(b"vollst\xc3\xa4ndige Handlung", landing)
        self.assertIn(b"Betriebliche Rollen statt beliebiger Gruppen", landing)
        self.assertIn(b"Geplante Einf\xc3\xbchrungspreise", landing)
        self.assertIn(b"8,90 \xe2\x82\xac", landing)
        self.assertIn(b"pro Monat", landing)
        self.assertIn(b"79 \xe2\x82\xac pro Jahr", landing)
        self.assertIn(b"Monatlich oder j\xc3\xa4hrlich buchbar", landing)
        self.assertIn(b"79 \xe2\x82\xac", landing)
        self.assertIn(b"299 \xe2\x82\xac", landing)
        self.assertIn(b"ab 599 \xe2\x82\xac", landing)
        self.assertIn(b'id="beta-dialog"', landing)
        self.assertIn(b"begrenzte Anzahl kostenfreier Beta-Testzug\xc3\xa4nge", landing)
        self.assertIn(b"14 Tage kostenfrei testen", landing)
        self.assertIn(b"f\xc3\xbcr 14 Tage", landing)
        self.assertNotIn(b"6\xe2\x80\x938 Wochen", landing)
        self.assertEqual(landing.count(b"Keine automatische Verl\xc3\xa4ngerung"), 1)
        self.assertIn(b"Nachhaltige Sch\xc3\xbclerfirma", landing)
        self.assertIn(b"Vom Projektauftrag bis zum gesicherten Ergebnis", landing)
        self.assertIn(b"VPS in Deutschland", landing)
        self.assertIn(b"Grunds\xc3\xa4tze der DSGVO", landing)
        self.assertIn(b'href="#sicherheit">Datenschutz</a>', landing)
        self.assertIn(b"Testzugang f\xc3\xbcr ProjektKontor anfragen", landing)
        self.assertIn(b"R\xc3\xbcckspr\xc3\xbcnge und \xc3\x9cberarbeitungen", landing)
        self.assertIn(b'class="process-lines"', landing)
        self.assertIn(b"L\xc3\xa4ngerfristiger<br>Unterrichtszusammenhang", landing)
        self.assertIn(b"Ergebnisse einer Einheit", landing)
        self.assertIn(b"\xc3\xbcbergeordnete Handlungsergebnis", landing)
        self.assertNotIn(b"Kein zus\xc3\xa4tzlicher Verwaltungsort", landing)
        self.assertIn(b'href="/login">Login', landing)
        self.assertIn(b'<link rel="canonical" href="https://projektkontor.org/">', landing)
        self.assertIn(b'<link rel="icon" type="image/png" sizes="64x64" href="/favicon.png">', landing)
        self.assertIn(b'<link rel="apple-touch-icon" sizes="180x180" href="/apple-touch-icon.png">', landing)
        self.assertNotIn(b'id="login-form"', landing)
        self.assertNotIn("X-Robots-Tag", headers)
        status, login, _ = self.client.request("GET", "/login")
        self.assertEqual(status, 200)
        self.assertIn(b'id="login-form"', login)
        self.assertIn(b'href="/">', login)
        self.assertIn(b"Zur Startseite", login)
        self.assertIn(b'content="noindex,nofollow,noarchive,nosnippet"', login)
        status, sitemap, _ = self.client.request("GET", "/sitemap.xml")
        self.assertEqual(status, 200)
        self.assertIn(b"https://projektkontor.org/", sitemap)
        status, impressum, _ = self.client.request("GET", "/impressum")
        self.assertEqual(status, 200)
        self.assertIn(b"Noch nicht f", impressum)
        self.assertIn(b'content="noindex,nofollow,noarchive"', impressum)
        self.assertIn(b"Name beziehungsweise Firma des Anbieters", impressum)
        self.assertIn(b"BIBB-Leitbild", landing)
        self.assertIn(b"Berufliche Handlungsf\xc3\xa4higkeit", landing)
        self.assertIn(b"Fachliche, methodische, soziale und pers\xc3\xb6nliche", landing)
        self.assertIn(b"Projektlernen", landing)
        self.assertIn(b"Selbststeuerung", landing)
        self.assertNotIn(b'<p class="card-index">A</p>', landing)
        status, landing_script, _ = self.client.request("GET", "/landing.js")
        self.assertEqual(status, 200)
        self.assertIn(b"target.closest('section')", landing_script)
        self.assertIn(b"availableHeight", landing_script)
        self.assertIn(b"largestVisibleArea", landing_script)
        self.assertIn(b"visibleArea", landing_script)
        self.assertIn(b"history.pushState", landing_script)
        self.assertIn(b"projektkontor-beta-popup-dismissed-v2", landing_script)
        self.assertEqual(landing_script.count(b"rememberBetaPopup();"), 1)
        self.assertIn(b"15000", landing_script)
        self.assertIn(b"Kostenloser Beta-Testzugang", landing_script)
        self.assertIn(b"scrollPosition", landing_script)
        self.assertIn(b"previousScrollBehavior", landing_script)

    def test_contact_form_is_human_checked_and_recipient_stays_server_side(self):
        payload = {
            "name": "Erika Beispiel", "email": "erika@example.org", "subject": "Testzugang",
            "message": "Ich möchte ProjektKontor gerne kennenlernen.", "privacy_accepted": True,
            "turnstile_token": "verified-token", "website": "",
        }
        with patch.object(self.app, "_verify_turnstile") as verify, patch.object(self.app, "_send_contact_email") as send:
            status, result, _ = self.client.request("POST", "/api/contact", payload)
        self.assertEqual(status, 200)
        self.assertTrue(result["ok"])
        verify.assert_called_once_with("verified-token", "unknown")
        send.assert_called_once_with("Erika Beispiel", "erika@example.org", "Testzugang", "Ich möchte ProjektKontor gerne kennenlernen.")
        saved = self.app.db.one("SELECT * FROM license_requests")
        self.assertEqual(saved["name"], "Erika Beispiel")
        self.assertEqual(saved["request_type"], "beta")
        self.assertEqual(saved["status"], "new")
        self.assertEqual(saved["email_status"], "sent")

        with patch.object(self.app, "_verify_turnstile"), patch.object(
            self.app, "_send_contact_email", side_effect=HttpError(503, "Mailversand nicht eingerichtet")
        ):
            status, result, _ = self.client.request(
                "POST", "/api/contact", {**payload, "email": "zweite@example.org"}
            )
        self.assertEqual(status, 200)
        self.assertTrue(result["ok"])
        failed = self.app.db.one(
            "SELECT * FROM license_requests WHERE email=?",
            ("zweite@example.org",),
        )
        self.assertEqual(failed["email_status"], "failed")

        admin = Client(self.app)
        admin.request("POST", "/api/setup", {
            "first_name": "Admin", "username": "verwaltung", "password": "sicheres-admin-kennwort"
        })
        status, requests, _ = admin.request("GET", "/api/license-requests")
        self.assertEqual(status, 200)
        self.assertEqual(len(requests), 2)
        status, updated, _ = admin.request(
            "PATCH", f"/api/license-requests/{requests[0]['id']}", {"status": "closed"}
        )
        self.assertEqual(status, 200)
        self.assertEqual(updated["status"], "closed")

        config = Config(
            "127.0.0.1", 8080, Path(self.temp.name), 25 * 1024 * 1024,
            smtp_host="smtp.example.org", smtp_username="sender@example.org",
            smtp_password="secret", smtp_sender="sender@example.org", contact_recipient="private@example.org",
        )
        mail_app = App(config)
        with patch("projektkontor.server.smtplib.SMTP_SSL") as smtp:
            smtp.return_value.__enter__.return_value = smtp.return_value
            mail_app._send_contact_email("Erika Beispiel", "erika@example.org", "Testzugang", "Nachricht")
            sent_mail = smtp.return_value.send_message.call_args.args[0]
        self.assertEqual(sent_mail["Subject"], "ProjektKontor Anfrage: Testzugang")
        self.assertEqual(sent_mail["To"], "private@example.org")
        self.assertEqual(sent_mail["Reply-To"], "Erika Beispiel <erika@example.org>")

    def test_first_login_requires_open_privacy_confirmation(self):
        client = Client(self.app, auto_privacy=False)
        status, pending, headers = client.request("POST", "/api/setup", {
            "first_name": "Jeroen", "username": "lehrkraft", "password": "sicheres-testkennwort"
        })
        self.assertEqual(status, 200)
        self.assertTrue(pending["privacy_required"])
        self.assertNotIn("Set-Cookie", headers)
        status, denied, _ = client.request("GET", "/api/classes")
        self.assertEqual(status, 401)
        self.assertIn("melden", denied["error"])
        status, information, _ = client.request("GET", "/api/privacy")
        self.assertEqual(status, 200)
        self.assertEqual(information["version"], pending["privacy_version"])
        self.assertIn("Verarbeitete Daten", [section["heading"] for section in information["sections"]])
        status, accepted, headers = client.request("POST", "/api/privacy/accept", {
            "privacy_token": pending["privacy_token"], "privacy_version": pending["privacy_version"]
        })
        self.assertEqual(status, 200)
        self.assertTrue(accepted["ok"])
        self.assertIn("Set-Cookie", headers)
        self.assertEqual(self.app.db.one("SELECT COUNT(*) count FROM privacy_acceptances")["count"], 1)
        status, _, _ = client.request("POST", "/api/privacy/accept", {
            "privacy_token": pending["privacy_token"], "privacy_version": pending["privacy_version"]
        })
        self.assertEqual(status, 401)
        client.request("POST", "/api/logout", {})
        status, logged_in, _ = client.request("POST", "/api/login", {
            "username": "lehrkraft", "password": "sicheres-testkennwort"
        })
        self.assertEqual(status, 200)
        self.assertTrue(logged_in["ok"])
        self.assertNotIn("privacy_required", logged_in)

    def test_students_cannot_read_other_team_contents(self):
        teacher = self.client
        self.setup_teacher(teacher)
        _, bootstrap, _ = teacher.request("GET", "/api/bootstrap")
        teacher_id = bootstrap["user"]["id"]
        _, klass, _ = teacher.request("POST", "/api/classes", {"name": "SI24"})
        _, imported, _ = teacher.request("POST", f"/api/classes/{klass['id']}/import", {
            "rows": [{"first_name": "Lena"}, {"first_name": "Noah"}]
        })
        lena, noah = imported["created"]
        _, project, _ = teacher.request("POST", "/api/projects", {
            "title": "Sicherheitsprojekt", "class_id": klass["id"], "has_start": False, "has_end": False,
            "project_lead_id": teacher_id, "member_ids": [lena["id"], noah["id"], teacher_id],
        })
        teacher.request("POST", f"/api/projects/{project['id']}/teams", {
            "name": "Einkauf", "color": "#2563EB", "member_ids": [lena["id"]], "lead_ids": []
        })
        _, team_b, _ = teacher.request("POST", f"/api/projects/{project['id']}/teams", {
            "name": "Vertrieb", "color": "#16A34A", "member_ids": [noah["id"]], "lead_ids": []
        })
        _, secret_task, _ = teacher.request("POST", f"/api/projects/{project['id']}/tasks", {
            "title": "Vertrauliche Teamaufgabe", "team_ids": [team_b["id"]]
        })
        pdf_data = base64.b64encode(b"%PDF-team-b").decode()
        _, secret_upload, _ = teacher.request("POST", f"/api/projects/{project['id']}/uploads", {
            "name": "vertrieb.pdf", "media_type": "application/pdf", "file_base64": pdf_data,
            "team_id": team_b["id"],
        })
        student = Client(self.app)
        student.request("POST", "/api/login", {"username": lena["username"], "password": lena["access_code"]})
        status, detail, _ = student.request("GET", f"/api/projects/{project['id']}")
        self.assertEqual(status, 200)
        self.assertNotIn(secret_task["ids"][0], [task["id"] for task in detail["tasks"]])
        self.assertNotIn(secret_upload["id"], [upload["id"] for upload in detail["uploads"]])
        status, _, _ = student.request("GET", f"/api/uploads/{secret_upload['id']}")
        self.assertEqual(status, 403)
        status, _, _ = student.request("POST", f"/api/tasks/{secret_task['ids'][0]}/comments", {"body": "Unzulässiger Zugriff"})
        self.assertEqual(status, 403)

    def test_existing_accounts_keep_login_available(self):
        self.setup_teacher()
        _, klass, _ = self.client.request("POST", "/api/classes", {"name": "LOGIN24"})
        _, student, _ = self.client.request("POST", f"/api/classes/{klass['id']}/users", {"first_name": "Lena"})
        self.client.request("POST", "/api/logout", {})

        status, bootstrap, _ = self.client.request("GET", "/api/bootstrap")
        self.assertEqual(status, 200)
        self.assertTrue(bootstrap["configured"])
        self.assertTrue(bootstrap["has_accounts"])
        self.assertFalse(bootstrap["setup_available"])
        self.assertIsNone(bootstrap["user"])

        status, logged_in, _ = self.client.request("POST", "/api/login", {
            "username": student["username"], "password": student["access_code"]
        })
        self.assertEqual(status, 200)
        self.assertTrue(logged_in["ok"])

        status, index, _ = self.client.request("GET", "/app")
        self.assertEqual(status, 200)
        self.assertIn(b"show-login", index)
        self.assertIn(b"show-setup", index)
        status, script, _ = self.client.request("GET", "/app.js")
        self.assertEqual(status, 200)
        self.assertIn(b"login_type", script)
        self.assertIn(b"W\xc3\xa4hlen Sie Ihren Zugang", script)

    def test_login_type_separates_student_teacher_and_admin(self):
        self.client.request("POST", "/api/setup", {
            "first_name": "Jeroen", "username": "verwaltung", "password": "sicheres-testkennwort"
        })
        _, teacher, _ = self.client.request("POST", "/api/teachers", {
            "first_name": "Frau Meyer", "username": "lehrkraft_meyer"
        })
        status, _, _ = self.client.request("POST", "/api/licenses", {
            "customer_name": "Testlizenz", "plan": "single", "seat_limit": 1,
            "amount_cents": 7900, "payment_status": "paid", "status": "active",
            "starts_on": "2020-01-01", "ends_on": "2099-12-31",
            "teacher_ids": [teacher["id"]],
        })
        self.assertEqual(status, 200)
        self.assertEqual(self.app.db.one("SELECT must_change_password FROM users WHERE id=?", (teacher["id"],))["must_change_password"], 1)
        teacher_client = Client(self.app)
        teacher_client.request("POST", "/api/login", {
            "login_type": "teacher", "username": "lehrkraft_meyer", "password": teacher["initial_password"]
        })
        _, klass, _ = teacher_client.request("POST", "/api/classes", {"name": "ROLLEN24"})
        _, student, _ = teacher_client.request("POST", f"/api/classes/{klass['id']}/users", {"first_name": "Lena"})
        self.assertEqual(self.app.db.one("SELECT must_change_password FROM users WHERE id=?", (teacher["id"],))["must_change_password"], 0)
        status, _, _ = self.client.request("PATCH", f"/api/teachers/{teacher['id']}", {
            "first_name": "Frau Meyer", "username": "lehrkraft_meyer", "active": True, "reset_password": True
        })
        self.assertEqual(status, 403)

        wrong_role = Client(self.app)
        status, _, _ = wrong_role.request("POST", "/api/login", {
            "login_type": "teacher", "username": "verwaltung", "password": "sicheres-testkennwort"
        })
        self.assertEqual(status, 401)
        status, _, _ = wrong_role.request("POST", "/api/login", {
            "login_type": "admin", "username": student["username"], "password": student["access_code"]
        })
        self.assertEqual(status, 401)

        admin_client = Client(self.app)
        status, _, _ = admin_client.request("POST", "/api/login", {
            "login_type": "admin", "username": "verwaltung", "password": "sicheres-testkennwort"
        })
        self.assertEqual(status, 200)
        status, index, _ = admin_client.request("GET", "/app")
        self.assertEqual(status, 200)
        self.assertIn(b'class="workspace-nav"', index)
        status, script, _ = admin_client.request("GET", "/app.js")
        self.assertEqual(status, 200)
        self.assertIn(b"state.user?.is_owner&&state.route!=='admin'", script)
        self.assertIn(b"$('.brand').href='#/admin/teachers'", script)
        _, admin_bootstrap, _ = admin_client.request("GET", "/api/bootstrap")
        self.assertEqual(admin_bootstrap["user"]["is_owner"], 1)

        another_teacher_client = Client(self.app)
        status, _, _ = another_teacher_client.request("POST", "/api/login", {
            "login_type": "teacher", "username": "lehrkraft_meyer", "password": TEST_TEACHER_PASSWORD
        })
        self.assertEqual(status, 200)

        student_client = Client(self.app)
        status, _, _ = student_client.request("POST", "/api/login", {
            "login_type": "student", "username": student["username"], "password": student["access_code"]
        })
        self.assertEqual(status, 200)

    def test_admin_setup_can_be_protected_by_private_token(self):
        root = Path(self.temp.name) / "protected-setup"
        root.mkdir()
        protected = App(Config(
            "127.0.0.1", 8080, root, 25 * 1024 * 1024,
            admin_setup_token="nur-fuer-den-betreiber",
        ))
        client = Client(protected)
        status, bootstrap, _ = client.request("GET", "/api/bootstrap")
        self.assertEqual(status, 200)
        self.assertTrue(bootstrap["setup_token_required"])
        status, _, _ = client.request("POST", "/api/setup", {
            "first_name": "Jeroen", "username": "admin",
            "password": "sicheres-admin-kennwort", "setup_token": "falsch",
        })
        self.assertEqual(status, 403)
        status, _, _ = client.request("POST", "/api/setup", {
            "first_name": "Jeroen", "username": "admin",
            "password": "sicheres-admin-kennwort",
            "setup_token": "nur-fuer-den-betreiber",
        })
        self.assertEqual(status, 200)
        account = protected.db.one("SELECT role,is_owner FROM users WHERE username='admin'")
        self.assertEqual(account, {"role": "teacher", "is_owner": 1})

    def test_legacy_owner_account_becomes_regular_teacher_without_password_change(self):
        self.client.request("POST", "/api/setup", {
            "first_name": "Jeroen", "username": "lehrkraft", "password": "sicheres-testkennwort"
        })
        self.client.request("POST", "/api/logout", {})
        self.app.db.execute("DELETE FROM schema_migrations WHERE version=2")
        self.app.db.initialize()
        account = self.app.db.one("SELECT username,is_owner FROM users WHERE username='lehrkraft'")
        self.assertEqual(account["is_owner"], 0)
        status, denied, _ = self.client.request("POST", "/api/login", {
            "login_type": "teacher", "username": "lehrkraft", "password": "sicheres-testkennwort"
        })
        self.assertEqual(status, 403)
        self.assertIn("noch nicht freigeschaltet", denied["error"])
        _, bootstrap, _ = self.client.request("GET", "/api/bootstrap")
        self.assertFalse(bootstrap["configured"])
        self.assertTrue(bootstrap["setup_available"])

    def test_admin_sees_teacher_data_only_with_support_code(self):
        admin = self.client
        admin.request("POST", "/api/setup", {
            "first_name": "Jeroen", "username": "verwaltung", "password": "sicheres-testkennwort"
        })
        status, teacher, _ = admin.request("POST", "/api/teachers", {
            "first_name": "Frau Meyer", "username": "lehrkraft_meyer"
        })
        self.assertEqual(status, 200)
        status, _, _ = admin.request("POST", "/api/licenses", {
            "customer_name": "Testlizenz", "plan": "single", "seat_limit": 1,
            "amount_cents": 7900, "payment_status": "paid", "status": "active",
            "starts_on": "2020-01-01", "ends_on": "2099-12-31",
            "teacher_ids": [teacher["id"]],
        })
        self.assertEqual(status, 200)
        teacher_client = Client(self.app)
        teacher_client.request("POST", "/api/login", {
            "login_type": "teacher", "username": "lehrkraft_meyer", "password": teacher["initial_password"]
        })
        _, klass, _ = teacher_client.request("POST", "/api/classes", {"name": "KL24A"})
        _, student, _ = teacher_client.request("POST", f"/api/classes/{klass['id']}/users", {"first_name": "Lena"})

        status, _, _ = admin.request("GET", f"/api/classes/{klass['id']}/users")
        self.assertEqual(status, 403)
        _, hidden_classes, _ = admin.request("GET", "/api/classes")
        self.assertEqual(hidden_classes, [])

        with patch.object(self.app, "_send_support_email") as send:
            status, support, _ = teacher_client.request("POST", "/api/support/requests", {
                "phone": "+49 170 1234567", "message": "Ich benötige Hilfe bei einer Schülerzuordnung."
            })
        self.assertEqual(status, 200)
        send.assert_called_once()
        status, redeemed, _ = admin.request("POST", "/api/support/redeem", {
            "support_code": support["support_code"]
        })
        self.assertEqual(status, 200)
        self.assertEqual(redeemed["teacher"]["username"], "lehrkraft_meyer")
        status, users, _ = admin.request("GET", f"/api/classes/{klass['id']}/users")
        self.assertEqual(status, 200)
        self.assertEqual(users[0]["id"], student["id"])

        status, _, _ = teacher_client.request("DELETE", f"/api/support/requests/{support['id']}", {})
        self.assertEqual(status, 200)
        status, _, _ = admin.request("GET", f"/api/classes/{klass['id']}/users")
        self.assertEqual(status, 403)


if __name__ == "__main__":
    unittest.main()
