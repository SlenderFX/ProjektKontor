from __future__ import annotations

import io
import base64
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .db import Database


NAVY = colors.HexColor("#0F2E5D")
SAGE = colors.HexColor("#6F8B74")
LIGHT = colors.HexColor("#E6E9ED")
PRIMEADVISORY_LOGO_B64 = Path(__file__).resolve().parent / "static" / "primeadvisory-logo.png.b64"


def safe_paragraph(value: Any) -> str:
    return escape(str(value or "")).replace("\n", "<br/>")


def money(cents: int) -> str:
    return f"{cents / 100:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") + " EUR"


def generate_invoice_pdf(invoice: dict[str, Any]) -> bytes:
    """Create an immutable, printable invoice from a database snapshot."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=f"Rechnung {invoice['invoice_number']}",
        author=invoice["issuer_name"],
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="InvoiceBrand", parent=styles["Heading1"], fontSize=19, leading=22,
        textColor=NAVY, spaceAfter=1 * mm,
    ))
    styles.add(ParagraphStyle(
        name="InvoiceRight", parent=styles["BodyText"], alignment=TA_RIGHT,
        fontSize=8.5, leading=11, textColor=colors.HexColor("#475467"),
    ))
    styles.add(ParagraphStyle(
        name="InvoiceSmall", parent=styles["BodyText"], fontSize=8, leading=10,
        textColor=colors.HexColor("#667085"),
    ))
    styles.add(ParagraphStyle(
        name="InvoiceTotal", parent=styles["BodyText"], fontSize=11, leading=14,
        textColor=NAVY, alignment=TA_RIGHT,
    ))

    issuer_lines = invoice["issuer_name"]
    if invoice.get("issuer_proprietor"):
        issuer_lines += f"<br/>{safe_paragraph(invoice['issuer_proprietor'])}"
    issuer_lines += f"<br/>{safe_paragraph(invoice['issuer_address'])}"
    if invoice.get("issuer_email"):
        issuer_lines += f"<br/>{safe_paragraph(invoice['issuer_email'])}"
    if PRIMEADVISORY_LOGO_B64.is_file():
        logo_bytes = base64.b64decode(PRIMEADVISORY_LOGO_B64.read_text(encoding="ascii"))
        brand = Image(io.BytesIO(logo_bytes), width=68 * mm, height=16 * mm)
    else:
        brand = Paragraph("PRIME<span color='#6F8B74'>Advisory</span>", styles["InvoiceBrand"])
    header = Table([
        [
            brand,
            Paragraph(issuer_lines, styles["InvoiceRight"]),
        ],
        [
            Paragraph("Rechnung für ProjektKontor", styles["InvoiceSmall"]),
            "",
        ],
    ], colWidths=[92 * mm, 78 * mm])
    header.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("LINEBELOW", (0, 1), (-1, 1), 1.2, SAGE),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 6 * mm),
    ]))

    recipient = invoice.get("organization", "").strip()
    if recipient:
        recipient += "<br/>"
    recipient += safe_paragraph(invoice["customer_name"])
    recipient += "<br/>" + safe_paragraph(invoice["billing_address"])
    details = [
        ["Rechnungsnummer", invoice["invoice_number"]],
        ["Rechnungsdatum", datetime.fromisoformat(invoice["issued_on"]).strftime("%d.%m.%Y")],
        ["Leistungsdatum", datetime.fromisoformat(invoice["service_on"]).strftime("%d.%m.%Y")],
        [
            "Zahlungsziel",
            "Sofort" if invoice["due_on"] == invoice["issued_on"]
            else datetime.fromisoformat(invoice["due_on"]).strftime("%d.%m.%Y"),
        ],
    ]
    if invoice.get("invoice_reference"):
        details.append(["Referenz", invoice["invoice_reference"]])
    detail_table = Table(details, colWidths=[30 * mm, 45 * mm])
    detail_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 0), (0, -1), NAVY),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    address_block = Table([
        [Paragraph("<b>Rechnung an</b><br/>" + recipient, styles["BodyText"]), detail_table]
    ], colWidths=[95 * mm, 75 * mm])
    address_block.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
    ]))

    item_rows = [
        [
            Paragraph("<b>Leistung</b>", styles["InvoiceSmall"]),
            Paragraph("<b>Zeitraum</b>", styles["InvoiceSmall"]),
            Paragraph("<b>Betrag</b>", styles["InvoiceSmall"]),
        ],
        [
            Paragraph(safe_paragraph(invoice["description"]), styles["BodyText"]),
            Paragraph(
                f"{datetime.fromisoformat(invoice['license_starts_on']).strftime('%d.%m.%Y')} bis "
                f"{datetime.fromisoformat(invoice['license_ends_on']).strftime('%d.%m.%Y')}",
                styles["InvoiceSmall"],
            ),
            Paragraph(money(invoice["net_cents"]), styles["InvoiceRight"]),
        ],
    ]
    items = Table(item_rows, colWidths=[90 * mm, 45 * mm, 35 * mm], repeatRows=1)
    items.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), LIGHT),
        ("TEXTCOLOR", (0, 0), (-1, 0), NAVY),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, SAGE),
        ("LINEBELOW", (0, 1), (-1, 1), 0.4, LIGHT),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (2, 0), (2, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    totals_rows = [["Zwischensumme", money(invoice["net_cents"])]]
    if invoice["vat_cents"]:
        totals_rows.append([
            f"Umsatzsteuer {invoice['vat_rate_basis_points'] / 100:.0f} %",
            money(invoice["vat_cents"]),
        ])
    totals_rows.append(["Gesamtbetrag", money(invoice["gross_cents"])])
    totals = Table(totals_rows, colWidths=[45 * mm, 35 * mm], hAlign="RIGHT")
    totals.setStyle(TableStyle([
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("TEXTCOLOR", (0, -1), (-1, -1), NAVY),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, -1), (-1, -1), 11),
        ("LINEABOVE", (0, -1), (-1, -1), 1.2, NAVY),
        ("TOPPADDING", (0, -1), (-1, -1), 7),
    ]))

    if invoice["gross_cents"] == 0:
        payment_parts = ["Nullrechnung: Für diese kostenfreie Leistung ist keine Zahlung erforderlich."]
    else:
        payment_parts = [
            (
                "Der Rechnungsbetrag ist sofort fällig."
                if invoice["due_on"] == invoice["issued_on"]
                else f"Bitte zahlen Sie den Gesamtbetrag bis zum {datetime.fromisoformat(invoice['due_on']).strftime('%d.%m.%Y')}."
            )
            + f" Verwenden Sie dabei die Rechnungsnummer {safe_paragraph(invoice['invoice_number'])}.",
        ]
        if invoice.get("iban"):
            payment_parts.append(f"IBAN: {safe_paragraph(invoice['iban'])}")
        if invoice.get("bic"):
            payment_parts.append(f"BIC: {safe_paragraph(invoice['bic'])}")
        if invoice.get("bank_name"):
            payment_parts.append(f"Bank: {safe_paragraph(invoice['bank_name'])}")

    footer_text = (
        f"{safe_paragraph(invoice['issuer_name'])}"
        + (f" · {safe_paragraph(invoice['issuer_proprietor'])}" if invoice.get("issuer_proprietor") else "")
        + f" · {safe_paragraph(invoice['issuer_address']).replace('<br/>', ' · ')}"
        + f"<br/>Steuerliche Kennung: {safe_paragraph(invoice['issuer_tax_identifier'])}"
    )
    story: list[Any] = [
        header,
        Spacer(1, 9 * mm),
        address_block,
        Spacer(1, 12 * mm),
        Paragraph("Rechnung", styles["Title"]),
        Paragraph(
            "Vielen Dank für Ihr Vertrauen in ProjektKontor. Wir berechnen die folgende Leistung:",
            styles["BodyText"],
        ),
        Spacer(1, 7 * mm),
        items,
        Spacer(1, 6 * mm),
        totals,
        Spacer(1, 8 * mm),
    ]
    if invoice.get("tax_note"):
        story += [Paragraph(safe_paragraph(invoice["tax_note"]), styles["InvoiceSmall"]), Spacer(1, 5 * mm)]
    story += [
        Paragraph("<br/>".join(payment_parts), styles["BodyText"]),
        Spacer(1, 15 * mm),
        Paragraph(footer_text, styles["InvoiceSmall"]),
    ]
    doc.build(story)
    return buffer.getvalue()


def report_date(value: str, include_time: bool) -> str:
    parsed = datetime.fromisoformat(value)
    return parsed.strftime("%d.%m.%Y, %H:%M Uhr") if include_time else parsed.strftime("%d.%m.%Y")


def normalize_username(first_name: str) -> str:
    value = first_name.strip().lower()
    value = value.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    value = re.sub(r"[^a-z0-9]", "", value)
    return value or "konto"


def next_username(connection, first_name: str) -> str:
    base = normalize_username(first_name)
    rows = connection.execute(
        "SELECT username FROM users WHERE username LIKE ? COLLATE NOCASE",
        (base + "__",),
    ).fetchall()
    used = {str(row[0]).lower() for row in rows}
    for number in range(1, 100):
        candidate = f"{base}{number:02d}"
        if candidate not in used:
            return candidate
    raise ValueError(f"Für {first_name} können keine weiteren Benutzernamen mit zwei Ziffern erzeugt werden.")


def account_template(default_class: str = "") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ProjektKontor Import"
    sheet.append(["Vorname", "Klasse"])
    sheet.append(["Lena", default_class or "GH24"])
    sheet.append(["Yusuf", default_class or "GH24"])
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 18
    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True, color="FFFFFF")
        cell.fill = cell.fill.copy(fill_type="solid", fgColor="0F2E5D")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def parse_account_workbook(data: bytes) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Die Excel-Datei konnte nicht gelesen werden.") from exc
    sheet = workbook.active
    header = [str(cell.value or "").strip().lower() for cell in next(sheet.iter_rows(max_row=1))]
    if "vorname" not in header:
        raise ValueError("Die Spalte „Vorname“ fehlt.")
    column = header.index("vorname")
    class_column = header.index("klasse") if "klasse" in header else None
    result: list[dict[str, Any]] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        value = str(row[column] or "").strip()
        class_name = str(row[class_column] or "").strip() if class_column is not None and class_column < len(row) else ""
        if not value:
            continue
        if len(value) > 80:
            result.append({"row": row_number, "first_name": value[:80], "class_name": class_name, "error": "Vorname ist zu lang."})
        else:
            result.append({"row": row_number, "first_name": value, "class_name": class_name, "error": ""})
    if not result:
        raise ValueError("Die Excel-Datei enthält keine Vornamen.")
    return result


def generate_project_report(db: Database, project_id: int, destination: Path) -> None:
    project = db.one(
        """SELECT p.*, c.name class_name, u.first_name lead_name
           FROM projects p JOIN classes c ON c.id=p.class_id
           JOIN users u ON u.id=p.project_lead_id WHERE p.id=?""",
        (project_id,),
    )
    if not project:
        raise ValueError("Projekt nicht gefunden")
    teams = db.all("SELECT * FROM teams WHERE project_id=? AND status='active' ORDER BY name", (project_id,))
    phases = db.all("SELECT * FROM phases WHERE project_id=? ORDER BY sort_order,id", (project_id,))
    tasks = db.all(
        """SELECT t.*, tm.name team_name, ph.name phase_name, s.label status_label
           FROM tasks t LEFT JOIN teams tm ON tm.id=t.team_id
           LEFT JOIN phases ph ON ph.id=t.phase_id
           LEFT JOIN task_statuses s ON s.project_id=t.project_id AND s.key=t.status_key
           WHERE t.project_id=? ORDER BY t.created_at""",
        (project_id,),
    )
    events = db.all(
        """SELECT e.*, u.first_name actor_name FROM project_events e
           LEFT JOIN users u ON u.id=e.actor_id WHERE e.project_id=? ORDER BY e.created_at""",
        (project_id,),
    )
    comments = db.all(
        """SELECT c.*, u.first_name author_name, t.title task_title FROM comments c
           JOIN users u ON u.id=c.author_id JOIN tasks t ON t.id=c.task_id
           WHERE t.project_id=? AND c.hidden_at IS NULL ORDER BY c.created_at""",
        (project_id,),
    )
    uploads = db.all("SELECT * FROM uploads WHERE project_id=? ORDER BY created_at", (project_id,))

    destination.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(destination), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="PKTitle", parent=styles["Title"], textColor=NAVY, alignment=TA_CENTER, spaceAfter=10*mm))
    styles.add(ParagraphStyle(name="PKH1", parent=styles["Heading1"], textColor=NAVY, spaceBefore=5*mm))
    styles.add(ParagraphStyle(name="PKSmall", parent=styles["BodyText"], fontSize=8, leading=10))
    story: list[Any] = [Paragraph("ProjektKontor", styles["PKTitle"]), Paragraph(safe_paragraph(project["title"]), styles["Title"])]
    story += [Spacer(1, 8*mm), Paragraph(safe_paragraph(project["description"] or "Ohne Projektbeschreibung"), styles["BodyText"])]
    story += [Spacer(1, 8*mm), Table([
        ["Klasse", project["class_name"]], ["Gesamtprojektleitung", project["lead_name"]],
        ["Zeitraum", (
            f"{report_date(project['start_at'], bool(project.get('start_has_time', 1)))} – {report_date(project['end_at'], bool(project.get('end_has_time', 1)))}" if project.get("has_start", 1) and project.get("has_end", 1)
            else f"ab {report_date(project['start_at'], bool(project.get('start_has_time', 1)))}" if project.get("has_start", 1)
            else f"bis {report_date(project['end_at'], bool(project.get('end_has_time', 1)))}" if project.get("has_end", 1)
            else "Nicht festgelegt"
        )], ["Bericht erstellt", datetime.now().strftime("%d.%m.%Y %H:%M")]
    ], colWidths=[50*mm, 110*mm], style=TableStyle([("BACKGROUND",(0,0),(0,-1),LIGHT),("TEXTCOLOR",(0,0),(0,-1),NAVY),("GRID",(0,0),(-1,-1),0.4,LIGHT),("VALIGN",(0,0),(-1,-1),"TOP")]))]

    story += [PageBreak(), Paragraph("Projektorganisation", styles["PKH1"])]
    for team in teams:
        members = db.all(
            """SELECT u.first_name, tm.is_lead, tm.business_role FROM team_members tm
               JOIN users u ON u.id=tm.user_id WHERE tm.team_id=? ORDER BY tm.is_lead DESC,u.first_name""",
            (team["id"],),
        )
        parts = []
        for member in members:
            label = member["first_name"]
            if member["is_lead"]:
                label += " (Teamleitung)"
            if member["business_role"]:
                label += f" – {member['business_role']}"
            parts.append(label)
        member_text = ", ".join(parts) or "Noch unbesetzt"
        story += [Paragraph(safe_paragraph(team["name"]), styles["Heading2"]), Paragraph(safe_paragraph(team["responsibility"] or "Keine Zuständigkeit beschrieben"), styles["BodyText"]), Paragraph(safe_paragraph(member_text), styles["PKSmall"])]

    story += [Paragraph("Projektphasen", styles["PKH1"])]
    if phases:
        phase_data = [["Phase", "Zeitraum", "Erwartetes Ergebnis"]] + [[p["name"], f"{p['start_at']} – {p['end_at']}", p["expected_result"]] for p in phases]
        story.append(Table(phase_data, repeatRows=1, colWidths=[42*mm,55*mm,63*mm], style=TableStyle([("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.4,LIGHT),("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),8)])))

    story += [Paragraph("Aufgaben und Ergebnisse", styles["PKH1"])]
    for task in tasks:
        assignees = db.all("SELECT u.first_name FROM task_assignees a JOIN users u ON u.id=a.user_id WHERE a.task_id=? ORDER BY u.first_name", (task["id"],))
        story += [Paragraph(safe_paragraph(task["title"]), styles["Heading2"]), Paragraph(safe_paragraph(
            f"Team: {task['team_name'] or 'ohne Team'} · Verantwortlich: {', '.join(a['first_name'] for a in assignees) or 'noch zu verteilen'} · Status: {task['status_label'] or task['status_key']}"), styles["PKSmall"])]
        if task["description"]:
            story.append(Paragraph(safe_paragraph(task["description"]), styles["BodyText"]))
        if task["result_text"]:
            story += [Paragraph("Ergebnis", styles["Heading3"]), Paragraph(safe_paragraph(task["result_text"]), styles["BodyText"])]

    story += [Paragraph("Kommentare", styles["PKH1"])]
    for comment in comments:
        edited = " · bearbeitet" if comment["edited_at"] else ""
        story.append(Paragraph(f"<b>{safe_paragraph(comment['author_name'])}</b> zu „{safe_paragraph(comment['task_title'])}“ · {safe_paragraph(comment['created_at'])}{edited}<br/>{safe_paragraph(comment['body'])}", styles["BodyText"]))

    story += [Paragraph("Projektverlauf", styles["PKH1"])]
    for event in events:
        story.append(Paragraph(safe_paragraph(f"{event['created_at']} · {event['actor_name'] or 'System'}: {event['message']}"), styles["PKSmall"]))

    story += [Paragraph("Dateiverzeichnis", styles["PKH1"])]
    if uploads:
        story.append(Table([["Datei","Typ","Größe"]] + [[u["original_name"],u["media_type"],f"{u['size_bytes']/1024:.0f} KB"] for u in uploads], repeatRows=1, colWidths=[90*mm,45*mm,25*mm], style=TableStyle([("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),0.4,LIGHT),("FONTSIZE",(0,0),(-1,-1),8)])))
        image_uploads=[upload for upload in uploads if str(upload["media_type"]).startswith("image/")]
        if image_uploads:
            story.append(Paragraph("Bildanhänge", styles["PKH1"]))
            for upload in image_uploads:
                image_path=db.path.parent/"uploads"/upload["stored_name"]
                if image_path.is_file():
                    try:
                        preview=Image(str(image_path));preview._restrictSize(155*mm,95*mm)
                        story += [Paragraph(safe_paragraph(upload["original_name"]),styles["Heading3"]),preview,Spacer(1,4*mm)]
                    except Exception:
                        story.append(Paragraph(safe_paragraph(f"{upload['original_name']} (Vorschau nicht verfügbar)"),styles["BodyText"]))
    else:
        story.append(Paragraph("Keine Uploads vorhanden.", styles["BodyText"]))

    doc.build(story)
